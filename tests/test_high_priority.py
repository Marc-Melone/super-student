"""Security and source-freshness regressions. Uses temporary libraries and dummy tokens only.

Run: python tests/test_high_priority.py
"""
from __future__ import annotations

import asyncio
import copy
import io
import os
import sys
import tempfile
import threading
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from unittest.mock import Mock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from superstudent import describe, index, mcp_server, notes, render
from superstudent.canvas import Canvas, DownloadError, _CanvasSession
from superstudent.config import DEFAULTS
from superstudent.extract import EXTRACTOR_VERSION
from superstudent.library import Library, LibraryPathError
from superstudent.outline import course_documents, write_outline
from superstudent.overview import write_course_files
from superstudent.packs import make_pack
from superstudent.sync import CourseSync, Syncer
from superstudent.util import atomic_write_json, parse_front_matter


class CredentialTests(unittest.TestCase):
    def test_download_temporary_file_does_not_follow_preexisting_link(self):
        with tempfile.TemporaryDirectory(prefix='ss-download-') as scratch:
            root = Path(scratch)
            dest, outside = root / 'lecture.txt', root / 'outside.txt'
            outside.write_text('PRIVATE_CONTENT')
            old_temp = root / f'.{dest.name}.{os.getpid()}.{threading.get_ident()}.part'
            old_temp.symlink_to(outside)
            response = Mock(headers={'Content-Type': 'text/plain'}, url='https://school.instructure.com/file')
            response.iter_content.return_value = [b'new course material']
            cv = Canvas('https://school.instructure.com', 'dummy')
            with patch.object(cv, '_get', return_value=response):
                cv._download_once(response.url, dest, max_bytes=100, expect_html=False, with_token=True)
            self.assertEqual(dest.read_bytes(), b'new course material')
            self.assertEqual(outside.read_text(), 'PRIVATE_CONTENT')
            self.assertTrue(old_temp.is_symlink())

    def test_authenticated_requests_require_exact_secure_origin(self):
        cv = Canvas('https://school.instructure.com', 'dummy-regression-token')
        session = Mock()
        session.get.return_value = Mock(status_code=200, headers={})
        with patch.object(cv, '_session', return_value=session):
            for url in ('http://school.instructure.com/files/1',
                        'https://school.instructure.com:8443/files/1',
                        'https://other.instructure.com/files/1',
                        'https://user@school.instructure.com/files/1',
                        'https://school.instructure.com:bad/files/1'):
                with self.subTest(url=url), self.assertRaises(ValueError):
                    cv._get(url)
            session.get.assert_not_called()
            cv._get('https://SCHOOL.instructure.com:443/files/1')
            self.assertEqual(session.get.call_args.kwargs['headers']['Authorization'], 'Bearer dummy-regression-token')

    def test_remote_http_and_embedded_credentials_are_rejected(self):
        for url in ('http://school.instructure.com', 'https://user:password@school.instructure.com'):
            with self.subTest(url=url), self.assertRaises(ValueError):
                Canvas(url, 'dummy')

    def test_redirect_strips_auth_on_any_origin_change(self):
        session = _CanvasSession()
        base = 'https://school.instructure.com'
        for url in ('http://school.instructure.com', 'https://school.instructure.com:8443', 'https://storage.example.com'):
            self.assertTrue(session.should_strip_auth(base, url))
        self.assertFalse(session.should_strip_auth(base, base + ':443'))
        self.assertFalse(session.should_strip_auth(base + '/a', base + '/b'))
        # Requests normally preserves auth on this upgrade; our exact-origin rule must override that.
        self.assertTrue(session.should_strip_auth('http://localhost', 'https://localhost'))

    def test_real_cross_port_redirect_does_not_forward_token(self):
        captured = []
        class Storage(BaseHTTPRequestHandler):
            def do_GET(self):
                captured.append(self.headers.get('Authorization'))
                self.send_response(200)
                self.end_headers()
                self.wfile.write(b'file content')
            def log_message(self, *args):
                pass
        storage = ThreadingHTTPServer(('127.0.0.1', 0), Storage)
        class Origin(BaseHTTPRequestHandler):
            def do_GET(self):
                captured.append(self.headers.get('Authorization'))
                self.send_response(302)
                self.send_header('Location', f'http://127.0.0.1:{storage.server_port}/file')
                self.end_headers()
            def log_message(self, *args):
                pass
        origin = ThreadingHTTPServer(('127.0.0.1', 0), Origin)
        threads = [threading.Thread(target=s.serve_forever, daemon=True) for s in (storage, origin)]
        for thread in threads:
            thread.start()
        try:
            cv = Canvas(f'http://127.0.0.1:{origin.server_port}', 'dummy-local-token')
            cv._get(cv.base + '/redirect').close()
            self.assertEqual(captured, ['Bearer dummy-local-token', None])
        finally:
            for server in (origin, storage):
                server.shutdown()
                server.server_close()
            for thread in threads:
                thread.join()


class LibraryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='ss-priority-')
        self.scratch = Path(self.temp.name)
        self.lib = Library(self.scratch / 'library')
        self.lib.ensure()
        self.folder = 'Fall/Course'
        self.course = self.lib.root / self.folder
        self.files = self.course / 'Files'
        self.files.mkdir(parents=True)
        self.lib.save_state({'courses': {'1': {'folder': self.folder, 'name': 'Course', 'code': 'C1', 'items': {}}}})
        self.outside = self.scratch / 'outside.md'
        self.outside.write_text('PRIVATE_OUTSIDE_MARKER')

    def tearDown(self):
        self.temp.cleanup()

    def rel(self, path):
        return path.relative_to(self.lib.root).as_posix()

    def picture(self, name='diagram.png'):
        from PIL import Image
        original = self.files / name
        Image.new('RGB', (80, 80), 'red').save(original)
        sidecar = original.with_name(original.name + '.md')
        sidecar.write_text('# Diagram\n\n## [Image]\n\n> Visual content: this file is an image. View it directly.\n')
        return original, sidecar

    def document(self, name='lesson.txt'):
        original = self.files / name
        original.write_text('First source version with important lecture facts.')
        sidecar = original.with_name(original.name + '.md')
        sidecar.write_text('# Lesson\n\n## [Page 1]\n\nImportant unique instructional content from the lecture.\n')
        return original, sidecar

    def syncer(self, original, sidecar):
        cv = Canvas('https://school.instructure.com', 'dummy-regression-token')
        syncer = Syncer(dict(DEFAULTS, library_dir=str(self.lib.root), transcribe='off'), cv, self.lib,
                        log=lambda _: None, media=False)
        cs = CourseSync(syncer, {'id': 1, 'name': 'Course', 'course_code': 'C1'})
        prev = {'path': original.relative_to(self.course).as_posix(), 'text': sidecar.relative_to(self.course).as_posix(),
                'stamp': 'old|50', 'status': 'ok', 'extractor': EXTRACTOR_VERSION,
                'last_successful_sync': '2026-10-01T00:00:00Z'}
        cs.items['file:7'] = copy.deepcopy(prev)
        return cs, cv, prev

    def call(self, name, args):
        with patch.object(mcp_server, '_lib', return_value=self.lib):
            return str(asyncio.run(mcp_server.build_server().call_tool(name, args)))

    def test_renderer_checks_adjacent_original(self):
        original, sidecar = self.picture()
        outside_image = self.scratch / 'outside.png'
        original.replace(outside_image)
        original.symlink_to(outside_image)
        with self.assertRaisesRegex(render.RenderError, 'outside the library'):
            render.render(self.lib, self.rel(sidecar))
        self.assertIsNone(render.page_count(self.lib, self.rel(sidecar)))
        self.assertNotIn(str(outside_image), self.call('view_page', {'path': self.rel(sidecar)}))

    def test_description_checks_adjacent_sidecar_and_metadata(self):
        original, sidecar = self.picture()
        sidecar.unlink()
        sidecar.symlink_to(self.outside)
        result = describe.save(self.lib, self.rel(original), 'Image', 'A sufficiently detailed diagram description.')
        self.assertFalse(result['ok'])
        self.assertEqual(self.outside.read_text(), 'PRIVATE_OUTSIDE_MARKER')
        sidecar.unlink()
        sidecar.write_text('# Diagram\n\n## [Image]\n\n> Visual content: this file is an image. View it directly.\n')
        (self.lib.meta / 'descriptions.json').symlink_to(self.outside)
        result = describe.save(self.lib, self.rel(original), 'Image', 'A sufficiently detailed diagram description.')
        self.assertFalse(result['ok'])
        self.assertNotIn('What this shows', sidecar.read_text())
        self.assertEqual(self.outside.read_text(), 'PRIVATE_OUTSIDE_MARKER')

    def test_note_sidecar_and_output_directory_are_contained(self):
        original, sidecar = self.document()
        sidecar.unlink()
        sidecar.symlink_to(self.outside)
        result = notes.save(self.lib, self.rel(original), 'Page 1: meaningful lecture notes. ' * 10)
        self.assertFalse(result['ok'])
        sidecar.unlink()
        sidecar.write_text('# Lesson\n\nMeaningful lecture content.\n')
        external_dir = self.scratch / 'outside-notes'
        external_dir.mkdir()
        (self.course / 'Study Notes').symlink_to(external_dir, target_is_directory=True)
        result = notes.save(self.lib, self.rel(original), 'Meaningful lecture notes. ' * 15)
        self.assertFalse(result['ok'])
        self.assertEqual(list(external_dir.iterdir()), [])

    def test_note_metadata_destination_checked_before_visible_write(self):
        original, _ = self.document()
        (self.lib.meta / 'notes.json').symlink_to(self.outside)
        result = notes.save(self.lib, self.rel(original), 'Page 1: meaningful lecture notes. ' * 10)
        self.assertFalse(result['ok'])
        self.assertFalse((self.course / 'Study Notes').exists())
        self.assertEqual(self.outside.read_text(), 'PRIVATE_OUTSIDE_MARKER')

    def test_mcp_generated_reads_check_every_final_file(self):
        for key, name in {'overview': 'COURSE_OVERVIEW.md', 'outline': 'OUTLINE.md', 'exam_intel': 'EXAM_INTEL.md',
                          'calendar': 'CALENDAR.md', 'grades': 'GRADES.md', 'links': 'LINKS.md',
                          'syllabus': 'Syllabus.md', 'notes': 'Study Notes/_Course notes.md',
                          'module:1': 'Modules/01 - Week/_Module Contents.md'}.items():
            target = self.course / name
            target.parent.mkdir(parents=True, exist_ok=True)
            target.symlink_to(self.outside)
            with self.subTest(file=key):
                response = self.call('course_file', {'course': 'C1', 'file': key})
                self.assertNotIn('PRIVATE_OUTSIDE_MARKER', response)
                self.assertIn('outside the library', response)
            target.unlink()

    def test_mcp_adjacent_read_and_write_tools_do_not_escape(self):
        original, sidecar = self.picture()
        sidecar.unlink()
        sidecar.symlink_to(self.outside)
        cases = [('read_material', {'path': self.rel(original)}),
                 ('save_visual_description', {'path': self.rel(original), 'where': 'Image', 'description': 'Detailed figure description. ' * 4}),
                 ('save_study_notes', {'path': self.rel(original), 'notes': 'Meaningful document notes. ' * 15})]
        for name, args in cases:
            with self.subTest(tool=name):
                self.assertNotIn('PRIVATE_OUTSIDE_MARKER', self.call(name, args))
        self.assertEqual(self.outside.read_text(), 'PRIVATE_OUTSIDE_MARKER')
        self.assertEqual(describe.survey(self.lib)['visuals'], 0)
        self.assertEqual(course_documents(self.lib, self.course), [])

    def test_gui_open_checks_adjacent_original(self):
        from superstudent.gui.app import App
        original, sidecar = self.picture()
        outside_image = self.scratch / 'outside.png'
        original.replace(outside_image)
        original.symlink_to(outside_image)
        app = App.__new__(App)
        with patch.object(app, 'lib', return_value=self.lib), patch.object(app, 'run_open') as opened:
            self.assertFalse(app.open_item(self.rel(sidecar))['ok'])
            opened.assert_not_called()

    def test_render_cache_and_index_companions_are_contained(self):
        from PIL import Image
        original, _ = self.picture()
        Image.new('RGB', (2000, 1800), 'blue').save(original)
        external_dir = self.scratch / 'outside-cache'
        external_dir.mkdir()
        self.lib.renders.symlink_to(external_dir, target_is_directory=True)
        with self.assertRaises(render.RenderError):
            render.render(self.lib, self.rel(original))
        self.assertEqual(list(external_dir.iterdir()), [])
        Path(str(self.lib.index_path) + '-wal').symlink_to(self.outside)
        with self.assertRaises(LibraryPathError):
            index.update_index(self.lib)
        self.assertEqual(self.outside.read_text(), 'PRIVATE_OUTSIDE_MARKER')

    def test_sync_outputs_and_extracted_assets_are_contained(self):
        original, sidecar = self.document()
        cs, cv, prev = self.syncer(original, sidecar)
        cv.download = Mock()
        sidecar.unlink()
        sidecar.symlink_to(self.outside)
        with self.assertRaises(LibraryPathError):
            cs._process_one('7', {'url': cv.base + '/file', 'updated_at': 'new', 'size': 60}, prev['path'], prev)
        cv.download.assert_not_called()
        sidecar.unlink()
        sidecar.write_text('# Safe document')
        external_dir = self.scratch / 'outside-assets'
        external_dir.mkdir()
        original.with_name(original.name + '.assets').symlink_to(external_dir, target_is_directory=True)
        with self.assertRaises(LibraryPathError):
            cs._process_one('7', {'url': cv.base + '/file'}, prev['path'], prev)
        self.assertEqual(list(external_dir.iterdir()), [])

    def _assert_my_files_asset_links_preserve_original_targets(self, extension, asset_name):
        from PIL import Image

        my_files = self.course / 'My Files'
        my_files.mkdir()
        original = my_files / ('lecture' + extension)
        picture = io.BytesIO()
        Image.new('RGB', (400, 400), 'red').save(picture, format='PNG')
        picture.seek(0)
        if extension == '.pptx':
            from pptx import Presentation
            from pptx.util import Inches

            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_picture(picture, Inches(1), Inches(1), Inches(2), Inches(2))
            presentation.save(original)
        else:
            from docx import Document
            from docx.shared import Inches

            document = Document()
            document.add_picture(picture, width=Inches(2))
            document.save(original)

        target = self.lib.root / 'existing-assets'
        target.mkdir()
        marker = target / 'keep.txt'
        marker.write_text('KEEP_INTERNAL_TARGET')
        linked_asset = target / asset_name
        linked_asset.symlink_to(self.outside)
        assets = original.with_name(original.name + '.assets')
        assets.symlink_to(target, target_is_directory=True)
        sidecar = original.with_name(original.name + '.md')
        cs, _, _ = self.syncer(original, sidecar)

        with patch('superstudent.extract._ocr.read_lines', return_value=[]):
            cs.process_my_files()

        self.assertEqual(self.outside.read_bytes(), b'PRIVATE_OUTSIDE_MARKER')
        self.assertEqual(marker.read_text(), 'KEEP_INTERNAL_TARGET')
        self.assertTrue(linked_asset.is_symlink())
        self.assertEqual(linked_asset.readlink(), self.outside)
        self.assertEqual({p.name for p in target.iterdir()}, {'keep.txt', asset_name})
        self.assertFalse(assets.is_symlink())
        self.assertEqual((assets / asset_name).read_bytes(), picture.getvalue())
        self.assertIn(asset_name, sidecar.read_text())

    def test_my_files_pptx_internal_assets_link_cannot_follow_external_child_link(self):
        self._assert_my_files_asset_links_preserve_original_targets('.pptx', 'slide-01-1.png')

    def test_my_files_docx_internal_assets_link_cannot_follow_external_child_link(self):
        self._assert_my_files_asset_links_preserve_original_targets('.docx', 'image-01.png')

    def test_outline_overview_and_pack_outputs_are_contained(self):
        (self.course / 'OUTLINE.md').symlink_to(self.outside)
        with self.assertRaises(LibraryPathError):
            write_outline(self.lib, self.course, 'Course')
        (self.course / 'GRADES.md').symlink_to(self.outside)
        with self.assertRaises(LibraryPathError):
            write_course_files(self.lib, '1', {'course': {}}, {}, self.course)
        (self.course / 'GRADES.md').unlink()
        pack_root = self.lib.root / '_Exam Packs'
        pack_root.symlink_to(self.scratch, target_is_directory=True)
        with self.assertRaises(LibraryPathError):
            make_pack(self.lib, self.folder)
        self.assertEqual(self.outside.read_text(), 'PRIVATE_OUTSIDE_MARKER')

    def test_office_conversion_checks_actual_output_before_converter_runs(self):
        source = self.files / 'Lecture.docx'
        source.write_bytes(b'fake Office source; converter will be mocked')
        with patch.object(render, 'office_pdf', return_value=None):
            self.assertIsNone(render._converted_pdf(self.lib, source))
        failed = next((self.lib.meta / 'converted').glob('*.failed'))
        conversion_dir = failed.with_suffix('')
        failed.unlink()
        conversion_dir.mkdir()
        (conversion_dir / 'Lecture.pdf').symlink_to(self.outside)
        with patch.object(render, 'office_pdf') as convert, self.assertRaises(LibraryPathError):
            render._converted_pdf(self.lib, source)
        convert.assert_not_called()
        self.assertEqual(self.outside.read_text(), 'PRIVATE_OUTSIDE_MARKER')

    def test_removed_my_files_cleanup_cannot_delete_external_files_or_directories(self):
        original, sidecar = self.document()
        cs, _, _ = self.syncer(original, sidecar)
        (self.course / 'My Files').mkdir()
        cs.items['mine:missing'] = {'text': str(self.outside), 'path': str(self.outside)}
        with self.assertRaises(LibraryPathError):
            cs.process_my_files()
        self.assertEqual(self.outside.read_text(), 'PRIVATE_OUTSIDE_MARKER')
        external_assets = self.scratch / 'victim.assets'
        external_assets.mkdir()
        marker = external_assets / 'private.txt'
        marker.write_text('KEEP')
        cs.items['mine:missing'] = {'text': 'My Files/gone.md', 'path': str(self.scratch / 'victim')}
        with self.assertRaises(LibraryPathError):
            cs.process_my_files()
        self.assertEqual(marker.read_text(), 'KEEP')

    def test_module_map_cleanup_cannot_follow_external_links(self):
        original, sidecar = self.document()
        cs, _, _ = self.syncer(original, sidecar)
        cs.listing_ok.add('modules')
        external_module = self.scratch / 'external-modules' / '01 - Old'
        external_module.mkdir(parents=True)
        contents = external_module / '_Module Contents.md'
        contents.write_text('KEEP')
        modules = self.course / 'Modules'
        modules.symlink_to(external_module.parent, target_is_directory=True)
        with self.assertRaises(LibraryPathError):
            cs.write_module_maps()
        self.assertEqual(contents.read_text(), 'KEEP')
        modules.unlink()
        local_module = modules / '01 - Old'
        local_module.mkdir(parents=True)
        (local_module / '_Module Contents.md').symlink_to(contents)
        with self.assertRaises(LibraryPathError):
            cs.write_module_maps()
        self.assertEqual(contents.read_text(), 'KEEP')

    def test_internal_links_still_work(self):
        original, _ = self.picture()
        alias = self.files / 'alias.png'
        alias.symlink_to(original)
        self.assertEqual(render.render(self.lib, self.rel(alias)), [original])

    def test_same_size_replacement_invalidates_descriptions_in_all_reads(self):
        original, sidecar = self.picture()
        original.write_bytes(b'oldimage')
        self.assertTrue(describe.save(self.lib, self.rel(original), 'Image', 'The old diagram shows obsolete purple anatomy.')['ok'])
        index.update_index(self.lib)
        self.assertTrue(index.search(self.lib, 'obsolete purple anatomy'))
        alias = self.files / 'alias.png.md'
        alias.symlink_to(sidecar)
        index.update_index(self.lib)
        timestamp = original.stat().st_mtime_ns
        original.write_bytes(b'newimage')
        os.utime(original, ns=(timestamp, timestamp))
        self.assertEqual(describe.current(self.lib, self.rel(original), original), {})
        self.assertEqual(describe.survey(self.lib)['described'], 0)
        self.assertNotIn('obsolete purple anatomy', index.read_document(self.lib, self.rel(sidecar)))
        self.assertEqual(index.search(self.lib, 'obsolete purple anatomy'), [])

    def test_summary_with_unavailable_sources_is_saved_as_incomplete(self):
        original, sidecar = self.document()
        cs, cv, prev = self.syncer(original, sidecar)
        cs._process_one('7', {'url': cv.base + '/file', 'locked_for_user': True}, prev['path'], prev)
        self.lib.save_state(cs.s.state)
        saved = notes.save(self.lib, self.folder, 'Course summary acknowledging source gaps. ' * 15)
        self.assertTrue(saved['ok'])
        self.assertIn('incomplete', saved['message'])
        self.assertIn('incomplete', notes.progress(self.lib)['courses'][0]['course_notes'])
        self.assertTrue(index.material_status(self.lib, saved['file'])['stale'])
        self.assertIn('Source warning', self.call('course_file', {'course': 'C1', 'file': 'notes'}))

    def test_short_encoded_and_asset_image_paths_retain_source_warnings(self):
        original, sidecar = self.picture('some diagram.png')
        cs, cv, prev = self.syncer(original, sidecar)
        cs._process_one('7', {'url': cv.base + '/file', 'locked_for_user': True}, prev['path'], prev)
        self.lib.save_state(cs.s.state)
        asset = original.with_name(original.name + '.assets') / 'figure.png'
        asset.parent.mkdir()
        asset.write_bytes(original.read_bytes())
        self.assertEqual(index.material_status(self.lib, self.rel(asset))['status'], 'restricted')
        for path in ('Files/some diagram.png', self.rel(original).replace(' ', '%20'), self.rel(asset)):
            with self.subTest(path=path):
                self.assertIn('Source warning', self.call('view_page', {'path': path}))

    def test_recording_replacement_pending_is_stale_until_success(self):
        original, sidecar = self.document('lecture.mp4')
        cs, cv, prev = self.syncer(original, sidecar)
        job = {'key': 'file:7', 'type': 'file', 'title': 'Lecture', 'rel_md': prev['text'], 'stamp': 'new|60'}
        cs.media_pending(job, 'network problem while downloading')
        self.lib.save_state(cs.s.state)
        self.assertEqual(index.material_status(self.lib, self.rel(sidecar))['status'], 'stale')
        self.assertIn('Source warning', index.read_document(self.lib, self.rel(sidecar)))
        self.assertEqual(notes.progress(self.lib)['courses'][0]['docs'], 0)
        from superstudent.media import Segment
        cs._write_transcript(job, [Segment(0, 1, 'Fresh replacement lecture facts.')], 'Canvas captions')
        self.lib.save_state(cs.s.state)
        self.assertEqual(index.material_status(self.lib, self.rel(sidecar))['status'], 'current')
        self.assertNotIn('Source warning', index.read_document(self.lib, self.rel(sidecar)))

    def test_unchanged_recording_newly_locked_bypasses_no_restriction(self):
        original, sidecar = self.document('lecture.mp4')
        cs, _, prev = self.syncer(original, sidecar)
        cs.file_meta['7'] = {'locked_for_user': True}
        job = {'key': 'file:7', 'type': 'file', 'file_id': '7', 'title': 'Lecture', 'rel_md': prev['text'], 'stamp': prev['stamp']}
        self.assertTrue(cs.media_up_to_date(job))
        self.lib.save_state(cs.s.state)
        self.assertEqual(index.material_status(self.lib, self.rel(sidecar))['status'], 'restricted')
        self.assertIn('historical copy', sidecar.read_text())

    def test_boolean_page_lock_invalidates_notes_and_study(self):
        page = self.course / 'Pages' / 'Locked.md'
        page.parent.mkdir()
        page.write_text('# Locked\n\nSome previous instructional text for this page.')
        state = self.lib.load_state()
        state['courses']['1']['items']['page:locked'] = {'path': 'Pages/Locked.md', 'locked': True}
        self.lib.save_state(state)
        self.assertEqual(index.material_status(self.lib, self.rel(page))['status'], 'restricted')
        self.assertEqual(notes.progress(self.lib)['courses'][0]['docs'], 0)

    def test_render_cache_changes_with_original_bytes(self):
        from PIL import Image
        original, _ = self.picture('diagram.bmp')
        first = render.render(self.lib, self.rel(original))[0]
        size, timestamp = original.stat().st_size, original.stat().st_mtime_ns
        Image.new('RGB', (80, 80), 'blue').save(original)
        self.assertEqual(original.stat().st_size, size)
        os.utime(original, ns=(timestamp, timestamp))
        second = render.render(self.lib, self.rel(original))[0]
        self.assertNotEqual(first, second)

    def test_legacy_size_only_descriptions_fail_closed(self):
        original, _ = self.picture()
        atomic_write_json(self.lib.meta / 'descriptions.json', {self.rel(original): {'Image': {'size': original.stat().st_size, 'text': 'Legacy text'}}})
        self.assertEqual(describe.current(self.lib, self.rel(original), original), {})

    def test_unchanged_file_newly_locked_is_labeled_and_excluded_from_study(self):
        original, sidecar = self.document()
        cs, cv, prev = self.syncer(original, sidecar)
        cs._process_one('7', {'url': cv.base + '/file', 'updated_at': 'old', 'size': 50, 'locked_for_user': True}, prev['path'], prev)
        self.lib.save_state(cs.s.state)
        self.assertEqual(cs.items['file:7']['status'], 'locked')
        self.assertEqual(cs.stats['files_unchanged'], 0)
        self.assertIn('historical copy', sidecar.read_text())
        index.update_index(self.lib)
        hits = index.search(self.lib, 'instructional content')
        self.assertEqual(hits[0]['status'], 'restricted')
        self.assertIn('SOURCE WARNING', mcp_server.format_hits('content', hits))
        for args in ({}, {'locator': 'Page 1'}, {'start': 2, 'max_chars': 30, 'lean': True}):
            self.assertIn('Source warning', index.read_document(self.lib, self.rel(original), **args))
        self.assertEqual(notes.progress(self.lib)['courses'][0]['docs'], 0)
        self.assertFalse(notes.save(self.lib, self.rel(original), 'Page 1 lecture notes. ' * 15)['ok'])

    def test_failed_replacement_preserves_old_version_and_recovers(self):
        original, sidecar = self.document()
        cs, cv, prev = self.syncer(original, sidecar)
        meta = {'url': cv.base + '/file', 'updated_at': 'new', 'size': 60}
        cv.download = Mock(side_effect=DownloadError('synthetic offline error'))
        cs._process_one('7', meta, prev['path'], prev)
        # Warnings must work before state is saved, including a sync interrupted at this point.
        self.assertTrue(index.material_status(self.lib, self.rel(original))['stale'])
        self.assertIn('Source warning', index.read_document(self.lib, self.rel(original), locator='Page 1'))
        self.lib.save_state(cs.s.state)
        item = cs.items['file:7']
        self.assertEqual(item['successful_stamp'], 'old|50')
        self.assertEqual(item['attempted_stamp'], 'new|60')
        self.assertEqual(item['last_successful_sync'], '2026-10-01T00:00:00Z')
        self.assertIn('instructional content', sidecar.read_text())
        index.update_index(self.lib)
        self.assertTrue(index.search(self.lib, 'instructional content')[0]['stale'])
        self.assertIn('Source warning', self.call('read_material', {'path': self.rel(original), 'at': 'Page 1'}))
        def download(url, dest, **kwargs):
            dest.write_text('Replacement source with fresh facts.')
            return {'content_type': 'text/plain'}
        cv.download = download
        cs._process_one('7', meta, prev['path'], copy.deepcopy(item))
        self.lib.save_state(cs.s.state)
        self.assertEqual(item['status'], 'ok')
        self.assertFalse(item.get('stale'))
        self.assertEqual(item['successful_stamp'], 'new|60')
        self.assertNotIn('Source warning', index.read_document(self.lib, self.rel(original)))
        self.assertNotIn('sync_status:', sidecar.read_text())

    def test_summary_and_document_notes_track_visual_original_changes(self):
        original, _ = self.document()
        self.assertTrue(notes.save(self.lib, self.rel(original), 'Page 1 lecture-specific notes. ' * 12)['ok'])
        self.assertTrue(notes.save(self.lib, self.folder, 'Course summary of lecture facts. ' * 16)['ok'])
        self.assertEqual(notes.progress(self.lib)['courses'][0]['course_notes'], 'done')
        original.write_text('Other source version with important lecture facts.')  # same length, unchanged sidecar
        progress = notes.progress(self.lib)['courses'][0]
        self.assertEqual(progress['changed'], 1)
        self.assertEqual(progress['course_notes'], 'changed since studied')
        self.assertIn('Source warning', self.call('course_file', {'course': 'C1', 'file': 'notes'}))
        record = notes.load(self.lib)[self.folder]
        self.assertIn('Source warning', index.read_document(self.lib, record['file']))

    def test_summary_invalidated_by_text_deletion_addition_and_restriction(self):
        original, sidecar = self.document()
        for change in ('text', 'delete', 'add', 'restrict'):
            with self.subTest(change=change):
                self.assertTrue(notes.save(self.lib, self.folder, 'Course summary of lecture facts. ' * 16)['ok'])
                if change == 'text':
                    sidecar.write_text(sidecar.read_text() + '\nNew course facts.')
                elif change == 'delete':
                    sidecar.unlink()
                elif change == 'add':
                    self.document('extra.txt')
                else:
                    state = self.lib.load_state()
                    state['courses']['1']['items']['file:2'] = {'path': 'Files/extra.txt', 'text': 'Files/extra.txt.md', 'status': 'locked'}
                    self.lib.save_state(state)
                self.assertEqual(notes.progress(self.lib)['courses'][0]['course_notes'], 'changed since studied')

    def test_module_summary_depends_on_current_source_set(self):
        module = self.course / 'Modules' / '01 - Week'
        module.mkdir(parents=True)
        source = module / 'Page.md'
        source.write_text('# Page\n\nInstructional source text for the module.')
        self.assertTrue(notes.save(self.lib, self.rel(module), 'Module summary of the reading. ' * 16)['ok'])
        self.assertEqual(notes.progress(self.lib)['courses'][0]['modules'][0]['status'], 'done')
        source.write_text(source.read_text() + '\nChanged facts for the module.')
        self.assertEqual(notes.progress(self.lib)['courses'][0]['modules'][0]['status'], 'changed since studied')

    def test_legacy_summary_requires_review(self):
        self.document()
        atomic_write_json(self.lib.meta / 'notes.json', {self.folder: {'kind': 'course', 'docs': 1, 'date': '2099-01-01'}})
        self.assertEqual(notes.progress(self.lib)['courses'][0]['course_notes'], 'changed since studied')


if __name__ == '__main__':
    unittest.main(verbosity=2)
