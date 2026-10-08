"""Tests for everything fixed after the final review (v1.6): converting files to text, the Canvas client,
token storage, and syncing over several runs (deleted/replaced content, a page unlocking, a renamed module,
a dropped connection, a revoked token).

Run:  python tests/test_review_fixes.py
"""

from __future__ import annotations

import json
import os
import random
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(HERE))

WORK = Path(os.environ.get("SS_TEST_DIR") or tempfile.mkdtemp(prefix="ss-fixes-")) / "fixes"
shutil.rmtree(WORK, ignore_errors=True)
WORK.mkdir(parents=True)
os.environ["SUPERSTUDENT_HOME"] = str(WORK / "app")
os.environ["SUPERSTUDENT_LIBRARY"] = str(WORK / "lib")
os.environ["CANVAS_TOKEN"] = "test-token-123"
os.environ["SUPERSTUDENT_APP_DRYRUN"] = "1"       # the app records what it would open instead of opening it

PASSED = []
SKIPPED = []
FONT = os.environ.get("SS_TEST_FONT") or ("/System/Library/Fonts/Supplemental/Arial.ttf" if sys.platform == "darwin"
                                         else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
if os.environ.get("SS_TEST_FONT"):
    from superstudent import slide_render
    slide_render.FONTS = {False: [FONT], True: [FONT]}  # deterministic drawing/OCR fixtures across platforms
    slide_render._font.cache_clear()


def check(cond, label, detail=""):
    if not cond:
        raise AssertionError(f"FAILED: {label} {detail}")
    PASSED.append(label)
    print(f"  ok  {label}")


def skip(label, why):
    SKIPPED.append(label)
    print(f"  --  {label} (skipped: {why})")


def read(p: Path) -> str:
    return p.read_text(encoding="utf-8")


# ============================================================== converters
def converters() -> None:
    print("\n== Converting to text")
    from superstudent.content import TOKEN_RE, block_editor_html, html_to_markdown, simple_html_to_markdown
    from superstudent.util import safe_name

    base = "https://school.instructure.com"
    md, refs = html_to_markdown(
        '<p>Read <a class="instructure_file_link" title="Lecture1.pdf" href="/courses/1/files/2?wrap=1" '
        'data-api-endpoint="https://school.instructure.com/api/v1/courses/1/files/2">Lecture1.pdf</a>, the '
        '<a title="Week 2" href="/courses/1/pages/week-2">week 2 page</a> and the '
        '<a href="/courses/999/pages/syllabus">other section\'s syllabus</a>.</p>', base, "1")
    tokens = [m.group(2) for m in TOKEN_RE.finditer(md)]
    check(tokens == ["file", "page"] and "2" in refs.files, "links with hover titles still become local links", md)
    check("https://school.instructure.com/courses/999/pages/syllabus" in md and "syllabus" not in refs.pages,
          "another course's page stays a Canvas link")
    check(refs.pages == {"week-2": "week 2 page"}, "linked pages are remembered so they can be fetched", refs.pages)
    md, _ = html_to_markdown("<p>E = mc<sup>2</sup>, 10<sup>-3</sup> m, x<sub>i+1</sub>, H<sub>2</sub>O</p>", base)
    check("mc^2" in md and "10^(-3)" in md and "x_(i+1)" in md and "H_2O" in md, "superscripts and subscripts kept", md)
    check("σ^2" in simple_html_to_markdown("<p>Var = &sigma;<sup>2</sup></p>"), "superscripts in Word files too")
    md, _ = html_to_markdown("<table><tr><th>Week</th><th>Topic</th><th>Exam</th></tr>"
                             "<tr><td rowspan='2'>1</td><td>Bonds</td><td colspan='1'>-</td></tr>"
                             "<tr><td>Yields</td><td>Quiz</td></tr></table>", base)
    rows = [r for r in md.splitlines() if r.startswith("|")]
    check(len(rows) == 4 and rows[3].count("|") == rows[0].count("|") and "| 1 |" in rows[3],
          "merged table cells stay under the right headings", md)
    html = block_editor_html({"blocks": json.dumps({
        "ROOT": {"type": {"resolvedName": "PageBlock"}, "nodes": ["a", "b"], "props": {}},
        "a": {"type": {"resolvedName": "HeadingBlock"}, "props": {"text": "Cell cycle", "level": 2}},
        "b": {"type": {"resolvedName": "ImageBlock"}, "props": {"src": "/courses/1/files/7/preview", "alt": "mitosis"}}})})
    md, refs = html_to_markdown(html, base, "1")
    check("## Cell cycle" in md and refs.files.get("7") == "mitosis", "block-editor pages keep text and pictures", md)
    long = "第三章 债券定价与收益率曲线的基本原理以及久期和凸性在利率风险管理中的应用 " * 3
    name = safe_name(long + ".pdf")
    check(len(name.encode()) <= 200 and name.endswith(".pdf"), f"long Chinese names fit the disk ({len(name.encode())} bytes)")
    (WORK / "names" / (name + ".assets")).mkdir(parents=True)
    (WORK / "names" / (name + ".md")).write_text("ok")
    check(True, "…including the .md and .assets that go with them")

    import pymupdf

    from superstudent import ocr
    from superstudent.extract import _csv, _docx, _image, _pdf, _pptx, _xlsx

    # Two-column handout: read down the left column, then the right.
    random.seed(1)
    paras = [f"PARAGRAPH {i} START. " + " ".join(f"word{i}x{j}" for j in range(random.randint(18, 30)))
             + f" PARAGRAPH {i} END." for i in range(1, 13)]
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    for rect, chunk in ((pymupdf.Rect(54, 54, 296, 738), paras[:6]), (pymupdf.Rect(316, 54, 558, 738), paras[6:])):
        assert page.insert_textbox(rect, "\n\n".join(chunk), fontsize=9) >= 0
    twocol = WORK / "twocol.pdf"
    doc.save(str(twocol))
    body = _pdf(twocol).body
    order = [int(n) for n in re.findall(r"PARAGRAPH (\d+) START", body)]
    check(order == list(range(1, 13)) and "PARAGRAPH 6 END" in body, "two-column PDF read column by column", order)

    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    for row, (a, b, c) in enumerate([("Week", "Topic", "Reading"), ("1", "Bond pricing", "Ch. 3"),
                                     ("2", "Duration", "Ch. 4")]):
        for x, text in ((72, a), (200, b), (420, c)):
            page.insert_text((x, 100 + row * 22), text, fontsize=11)
    table_pdf = WORK / "schedule.pdf"
    doc.save(str(table_pdf))
    body = _pdf(table_pdf).body
    check("Week | Topic | Reading" in body and "2 | Duration | Ch. 4" in body, "table rows in a PDF stay rows", body[-300:])

    if ocr.engine():
        src = pymupdf.open(str(twocol))
        scan = pymupdf.open()
        sp = scan.new_page(width=src[0].rect.width, height=src[0].rect.height)
        sp.insert_image(sp.rect, pixmap=src[0].get_pixmap(dpi=150))
        scanned = WORK / "scanned2col.pdf"
        scan.save(str(scanned))
        order = [int(n) for n in re.findall(r"PARAGRAPH (\d+) START", _pdf(scanned).body)]
        check(order == list(range(1, 13)), "scanned two-column page read in column order", order)
    else:
        skip("scanned two-column page read in column order", "no OCR engine here")

    from PIL import Image, ImageDraw, ImageFont

    font = ImageFont.truetype(FONT, 34) if Path(FONT).exists() else None
    if ocr.engine() and font:
        upright = Image.new("RGB", (900, 1200), "white")
        d = ImageDraw.Draw(upright)
        for i, line in enumerate(["Midterm review sheet", "Duration measures price sensitivity"]):
            d.text((60, 120 + i * 90), line, fill="black", font=font)
        exif = Image.Exif()
        exif[0x0112] = 6                       # the camera stored it sideways; viewers rotate it
        photo = WORK / "phone_photo.jpg"
        upright.rotate(90, expand=True).save(photo, exif=exif.tobytes(), quality=92)
        text = _image(photo).body
        check("Midterm review sheet" in text, "sideways phone photo read upright", text[-200:])
    else:
        skip("sideways phone photo read upright", "no OCR engine or font")

    from lxml import etree
    from pptx import Presentation
    from pptx.chart.data import XyChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Bond pricing"
    s.shapes._spTree.append(etree.fromstring(EQUATION_SHAPE))
    eq = WORK / "eq.pptx"
    prs.save(str(eq))
    body = _pptx(eq).body
    check("P=∑_(t=1)^T C/((1+y)^t)" in body and "PV of its cash flows" in body, "equations in slides kept as text", body)
    if ocr.engine():
        from superstudent.slide_render import SlidePainter

        deck = Presentation(str(eq))
        SlidePainter(deck, deck.slides[0]).paint().save(WORK / "eq-slide.png")
        drawn = ocr.read_text(WORK / "eq-slide.png")
        check("cash flows" in drawn and "P=" in drawn.replace(" ", ""), "…and drawn on the slide picture", drawn)
    else:
        skip("…and drawn on the slide picture", "no OCR engine here")

    from superstudent.extract import office_pdf, soffice_path

    if soffice_path():
        prs = Presentation()
        for i in range(1, 6):
            s = prs.slides.add_slide(prs.slide_layouts[5])
            s.shapes.title.text = f"Deck slide {i}" if i != 2 else "Hidden slide two"
        prs.slides[1]._element.set("show", "0")
        hidden = WORK / "hidden.pptx"
        prs.save(str(hidden))
        pdf = office_pdf(hidden, WORK / "hidden-pdf")
        doc = pymupdf.open(str(pdf)) if pdf else None
        check(doc is not None and doc.page_count == 5 and "Hidden slide two" in doc[1].get_text()
              and "Deck slide 3" in doc[2].get_text(), "with LibreOffice, slide N's picture is slide N (hidden slides kept)")
    else:
        skip("with LibreOffice, slide N's picture is slide N", "LibreOffice isn't installed here")

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Yield curve"
    data = XyChartData()
    series = data.add_series("Yield")
    for x, y in [(1, 4.1), (2, 4.3), (5, 4.6), (10, 4.9)]:
        series.add_data_point(x, y)
    s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(1), Inches(1.5), Inches(6), Inches(4), data)
    xy = WORK / "xy.pptx"
    prs.save(str(xy))
    check("| Yield | 10 | 4.9 |" in _pptx(xy).body, "scatter chart points kept")

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    for c in range(1, 41):
        ws.cell(row=1, column=c, value=f"Y{c}")
        ws.cell(row=2, column=c, value=c * 100)
    wide = WORK / "wide.xlsx"
    wb.save(wide)
    check("Y40" in _xlsx(wide).body, "40-column spreadsheet kept whole")
    wide_csv = WORK / "wide.csv"
    wide_csv.write_text(",".join(f"c{i}" for i in range(40)) + "\n" + ",".join(str(i) for i in range(40)) + "\n")
    check("c39" in _csv(wide_csv).body, "wide CSV kept whole")

    from docx import Document
    from docx.shared import Inches as DocInches

    d = Document()
    p = d.add_paragraph("Present value: ")
    p._p.append(etree.fromstring(
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>PV=</m:t></m:r>'
        '<m:f><m:num><m:r><m:t>C</m:t></m:r></m:num><m:den><m:sSup><m:e><m:r><m:t>(1+r)</m:t></m:r></m:e>'
        '<m:sup><m:r><m:t>n</m:t></m:r></m:sup></m:sSup></m:den></m:f></m:oMath>'))
    p = d.add_paragraph("Variance is σ")
    p.add_run("2").font.superscript = True
    p.add_run(" and the gas is CO")
    p.add_run("2").font.subscript = True
    math_docx = WORK / "math.docx"
    d.save(str(math_docx))
    body = _docx(math_docx).body
    check("[equation: PV=C/((1+r)^n)]" in body, "Word equations kept as text", body)
    check("σ^2" in body and "CO_2" in body, "Word superscripts and subscripts kept", body)

    if font:
        img = Image.new("RGB", (300, 60), "white")
        ImageDraw.Draw(img).text((10, 15), "P = C / (1+y)^t", fill="black", font=ImageFont.truetype(FONT, 24))
        img.save(WORK / "eq.png")
        d = Document()
        d.add_paragraph("The bond price is:")
        d.add_picture(str(WORK / "eq.png"), width=DocInches(2.5))
        small = WORK / "small_eq.docx"
        d.save(str(small))
        res = _docx(small)
        check("image-01.png" in res.body and res.visual_units, "small formula picture in a Word file kept", res.body)
    else:
        skip("small formula picture in a Word file kept", "no font")


EQUATION_SHAPE = '''<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
  xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
 <mc:Choice xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main" Requires="a14">
  <p:sp><p:nvSpPr><p:cNvPr id="10" name="TextBox 9"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>
   <p:spPr><a:xfrm><a:off x="914400" y="3200400"/><a:ext cx="7315200" cy="914400"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>
   <p:txBody><a:bodyPr/><a:lstStyle/>
    <a:p><a:r><a:rPr lang="en-US"/><a:t>The price of a bond is the PV of its cash flows, where y is the yield:</a:t></a:r></a:p>
    <a:p><a14:m><m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:oMath>
      <m:r><m:t>P=</m:t></m:r><m:nary><m:naryPr><m:chr m:val="&#8721;"/></m:naryPr><m:sub><m:r><m:t>t=1</m:t></m:r></m:sub><m:sup><m:r><m:t>T</m:t></m:r></m:sup>
      <m:e><m:f><m:num><m:r><m:t>C</m:t></m:r></m:num><m:den><m:sSup><m:e><m:r><m:t>(1+y)</m:t></m:r></m:e><m:sup><m:r><m:t>t</m:t></m:r></m:sup></m:sSup></m:den></m:f></m:e></m:nary>
    </m:oMath></m:oMathPara></a14:m></a:p>
   </p:txBody></p:sp>
 </mc:Choice>
 <mc:Fallback>
  <p:sp><p:nvSpPr><p:cNvPr id="10" name="TextBox 9"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>
   <p:spPr><a:xfrm><a:off x="914400" y="3200400"/><a:ext cx="7315200" cy="914400"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>
   <p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US"/><a:t> </a:t></a:r></a:p></p:txBody></p:sp>
 </mc:Fallback>
</mc:AlternateContent>'''


# ============================================================== Canvas client
def client(fake) -> None:
    print("\n== Talking to Canvas")
    from fake_canvas import TOKEN
    from superstudent.canvas import AuthError, Canvas, CanvasError, ForbiddenError, _without_verifier

    cv = Canvas(fake.base, TOKEN)
    try:
        cv.get("/api/v1/courses/101/media_attachments")
        got = None
    except CanvasError as exc:
        got = exc
    check(isinstance(got, ForbiddenError) and not isinstance(got, AuthError) and cv.auth_failed is None,
          "'not authorized' for one item doesn't count as a bad token", repr(got))
    bad = Canvas(fake.base, "wrong-token")
    try:
        bad.get("/api/v1/users/self")
        got = None
    except CanvasError as exc:
        got = exc
    check(isinstance(got, AuthError) and bad.auth_failed is got, "a rejected token is recognized")

    class Resp:
        def __init__(self, body: bytes):
            self.content = body
            self.headers = {}

        def json(self):
            return json.loads(self.content)

    replies = [Resp(b'{"id": 4'), Resp(b'{"id": 42}')]
    flaky = Canvas(fake.base, TOKEN)
    flaky._get = lambda url, params=None, **kw: replies.pop(0)
    flaky._sleep = staticmethod(lambda *a, **k: None)
    check(flaky.get("/api/v1/users/self") == {"id": 42}, "a cut-off answer from Canvas is asked for again")
    check(_without_verifier("https://s.edu/files/1/download?download_frd=1&verifier=abc") ==
          "https://s.edu/files/1/download?download_frd=1", "retired verifier links are not relied on")

    target = WORK / "dl" / "notes.pdf"
    fake.flaky_once = {1002}
    info = Canvas(fake.base, TOKEN).download(f"{fake.base}/files/1002/download?download_frd=1&verifier=v1002",
                                             target, file_id="1002")
    check(target.exists() and info["size"] == target.stat().st_size and not list(target.parent.glob(".*.part")),
          "a download cut off midway is retried, with no partial file left")
    check(fake.token_on_storage == 0, "the token never reaches the file-storage host")


# ============================================================== token storage
def keychain() -> None:
    print("\n== Saving the token")
    import superstudent.config as config

    calls = []
    stored = {}

    def fake_run(args, input=None, **kw):
        calls.append((list(args), input))
        if args[:2] == ["security", "-i"]:
            m = re.search(r"-w (\S+)", input or "")
            stored["token"] = m.group(1) if m else None
            return subprocess.CompletedProcess(args, 0, "", "")
        if "find-generic-password" in args:
            return subprocess.CompletedProcess(args, 0 if stored.get("token") else 44, (stored.get("token") or "") + "\n", "")
        return subprocess.CompletedProcess(args, 0, "", "")

    saved = (config.subprocess.run, config._keychain_available, config.TOKEN_FILE)
    config.subprocess.run = fake_run
    config._keychain_available = lambda: True
    config.TOKEN_FILE = WORK / "token-file"
    try:
        where = config.set_token({"canvas_url": "https://school.instructure.com"}, "7~SecretTokenValue123")
    finally:
        config.subprocess.run, config._keychain_available, config.TOKEN_FILE = saved
    on_command_line = any("7~SecretTokenValue123" in " ".join(args) for args, _ in calls)
    check(where == "your macOS Keychain" and not on_command_line and stored.get("token") == "7~SecretTokenValue123",
          "token saved to the Keychain without appearing on a command line")


# ============================================================== syncing
def syncing(fake, fixtures) -> None:
    import superstudent.media as media
    from superstudent.media import Segment

    media.detect_backend = lambda preference="auto": "faster-whisper"
    media.Transcriber.transcribe = lambda self, path, prompt="": [
        Segment(0.5, 6.0, "Welcome back. Today is all about modified duration.")]
    from superstudent import describe, notes
    from superstudent.cli import main as cli
    from superstudent.index import search
    from superstudent.library import Library

    lib = Library(WORK / "lib")
    fin = lib.root / "Fall 2026" / "FIN 6100 - Fixed Income"
    econ = lib.root / "Fall 2026" / "ECON 5200 - Econometrics"
    week2 = fin / "Modules" / "02 - Week 2 - Duration"
    week3 = fin / "Modules" / "03 - Week 3 - Immunization"
    removed = fin / "_Removed from Canvas"

    def stats(course: str) -> dict:
        last = json.loads(lib.last_sync_path.read_text())
        return next(c["stats"] for c in last["courses"] if course in c["name"])

    print("\n== Sync 1: everything a student can see")
    fake.extras = True
    fake.broken_downloads = {1005}
    fake.storage_hits.clear()
    check(cli(["sync", "--quiet"]) == 0, "sync finishes")
    page = read(week3 / "Week 3 Overview.md")
    check("## Immunization" in page and "main idea for the final" in page and "Practice problems" in page,
          "page made with the block editor has its text")
    check("_Locked: This page is locked until Oct 5._" in read(week3 / "Week 3 Preview.md"), "locked page says so")
    proof = econ / "Pages" / "Gauss-Markov proof.md"
    check(proof.exists() and "best linear unbiased" in read(proof), "page reachable only through a link is fetched")
    check("Gauss-Markov%20proof.md" in read(econ / "Modules" / "01 - Unit 1 Regression" / "OLS assumptions.md"),
          "…and the link points to it")
    quiz = read(fin / "Quizzes" / "Quiz 1 Bond Pricing.md")
    check("## Question review" in quiz and "- fall **(correct)**" in quiz and "**Your answer:** fall" in quiz
          and "4.31 ± 0.05 **(correct)**" in quiz and "**Your answer:** 4.2" in quiz, "past quiz questions with answers")
    ps1 = read(fin / "Assignments" / "Problem Set 1.md")
    row = next((ln for ln in ps1.splitlines() if ln.startswith("| Correct pricing")), "")
    check("semiannual compounding / annual is wrong" in row and "Partial (5): One bond priced incorrectly" in row
          and row.count("|") == 4, "rubric details kept on one table row", row)
    feedback = fin / "Assignments" / "Feedback" / "Problem Set 1"
    check((feedback / "PS1 marked up.pdf").exists() and "PS1%20marked%20up.pdf" in ps1, "instructor's marked-up file synced")
    audio = feedback / "Audio feedback on Problem Set 1 (recording).md"
    check(audio.exists() and "modified duration" in read(audio), "recorded audio feedback transcribed")
    check((fin / "Assignments" / "My Submissions" / "Problem Set 1" / "PS1 Marc.docx").exists(),
          "submitted file fetched through its signed link")
    thread = read(fin / "Discussions" / "Duration intuition.md")
    check("- _(deleted reply)_\n  - **Prof. Rivera (Instructor)**" in thread, "replies under a deleted post kept")
    check("    - Expect a duration question on the final." in thread, "multi-line replies keep their formatting")
    intel = read(fin / "EXAM_INTEL.md")
    check("Expect a duration question on the final." in intel, "instructor's whole reply reaches exam intel")
    check("only cover chapters 1 and 2" not in intel and (fin / "Discussions" / "Midterm rumor.md").exists(),
          "a classmate's exam rumor isn't reported as an exam hint")
    cal = read(fin / "CALENDAR.md")
    check("Discussion: yield curves: Reply to topic" in cal and "Required replies" in cal, "discussion checkpoint deadlines")
    lecture = read(week2 / "Week 2 lecture (lecture video).md")
    check("Canvas captions" in lecture and "make sure you know how to compute modified duration" in lecture
          and "couldn't be downloaded" in lecture, "captions kept when the video can't be downloaded")
    check((fin / "Files" / "Exams" / "SSO and SAML basics.pdf").exists(), "a PDF named like a sign-in address isn't rejected")
    check("sign-in page" in read(fin / "Files" / "Exams" / "Login trap.pdf.md"), "a real sign-in page still is")
    check(any(p.name.endswith(".pdf") and len(p.name.encode()) <= 200 for p in (fin / "Files" / "Exams").iterdir()
              if p.name.startswith("第三章")), "file with a very long Chinese name saved")
    check(fake.token_on_storage == 0 and fake.storage_hits, "token never sent to file storage")

    print("\n== Sync 2: the instructor replaces a file and unlocks a page; a download drops once")
    notes_pdf = f"{fin.relative_to(lib.root).as_posix()}/Modules/02 - Week 2 - Duration/Duration Notes.pdf"
    r = describe.save(lib, notes_pdf, "Page 2", "A price-yield curve: price falls as yield rises, curving toward the axis.", by="Claude")
    check(r["ok"], "picture description saved", r)
    r = notes.save(lib, notes_pdf, "Page 1: duration is the weighted average time of the cash flows. Page 2: the price-yield "
                   "curve and convexity. Page 3: worked example of modified duration, D_mod = D/(1+y).", by="Claude")
    check(r["ok"], "study notes saved", r)
    fake.replace_duration_notes = True
    fake.page_locked = False
    fake.flaky_once = {1012}
    check(cli(["sync", "--quiet"]) == 0, "sync finishes")
    check("immunization and the key-rate durations" in read(week3 / "Week 3 Preview.md"), "page synced while locked fills in once unlocked")
    gone = removed / "Modules" / "02 - Week 2 - Duration" / "Duration Notes.pdf"
    text = read(Path(str(gone) + ".md"))
    check(gone.exists() and "> **Removed from Canvas** on" in text and "removed_from_canvas:" in text,
          "deleted file moved to '_Removed from Canvas' and marked")
    replacement = [p for p in week2.iterdir() if p.name.startswith("Duration Notes") and p.suffix == ".pdf"]
    check(len(replacement) == 1 and not list(lib.root.rglob(".*.part")), "replacement downloaded despite a dropped connection")
    descs = json.loads((lib.meta / "descriptions.json").read_text())
    check(any("_Removed from Canvas" in k for k in descs) and "> **What this shows**" in text, "picture description moved with it")
    check(any("_Removed from Canvas" in k for k in json.loads((lib.meta / "notes.json").read_text())), "study notes moved with it")
    hits = search(lib, "modified duration", course="FIN", limit=40)
    flags = [h["removed"] for h in hits]
    check(True in flags and flags == sorted(flags), "search lists current material first, removed copies after", flags)
    check("_Removed" not in read(fin / "EXAM_INTEL.md"), "exam intel ignores removed material")
    check("**No longer on Canvas:** 1 item" in read(fin / "COURSE_OVERVIEW.md"), "overview says what was removed")
    check(stats("FIN").get("removed_from_canvas") == 1, "exactly one item retired")

    print("\n== Sync 3: the module is renamed")
    r = notes.save(lib, f"{fin.relative_to(lib.root).as_posix()}/Modules/02 - Week 2 - Duration",
                   "Week 2 ties duration to price risk. " * 12, by="Claude")
    check(r["ok"], "module notes saved", r)
    lecture_md = f"{fin.relative_to(lib.root).as_posix()}/Modules/02 - Week 2 - Duration/Week 2 Lecture Recording.mp4.md"
    r = notes.save(lib, lecture_md, "The lecture opens with modified duration, then walks through the convexity "
                   "adjustment and a worked example on the board; the professor stresses the convexity adjustment.", by="Claude")
    check(r["ok"], "lecture notes saved", r)
    fake.rename_week2 = True
    check(cli(["sync", "--quiet"]) == 0, "sync finishes")
    new_week2 = fin / "Modules" / "02 - Week 2 - Duration and Convexity"
    check(not week2.exists(), "old module folder and its stale contents list are gone")
    check((new_week2 / "Duration Notes.pdf").exists() and (new_week2 / "_Module Contents.md").exists(),
          "replacement now has the plain name in the renamed module")
    module_notes = fin / "Study Notes" / "Modules" / "02 - Week 2 - Duration and Convexity" / "_Module notes.md"
    check(module_notes.exists(), "module notes follow the renamed module")
    keys = json.loads((lib.meta / "notes.json").read_text())
    check(lecture_md.replace("02 - Week 2 - Duration", "02 - Week 2 - Duration and Convexity") in keys,
          "lecture notes follow the recording's transcript")
    check("Duration%20Notes.pdf" in read(new_week2 / "_Module Contents.md"), "module contents link the new file")

    print("\n== Sync 4: nothing changed")
    check(cli(["sync", "--quiet"]) == 0, "sync finishes")
    for course in ("FIN", "ECON"):
        st = stats(course)
        check(not st.get("docs_written") and not st.get("files_downloaded") and not st.get("files_extracted"),
              f"{course}: nothing rewritten or downloaded again", st)

    print("\n== Sync 5: the token is revoked partway through")
    before = sorted(p.relative_to(lib.root).as_posix() for p in lib.root.rglob("*") if "_Removed" in str(p))
    state_before = json.loads(lib.state_path.read_text())
    fake.revoke_after = 6
    fake.api_calls = 0
    check(cli(["sync", "--quiet"]) == 2, "sync stops with a sign-in message")
    after = sorted(p.relative_to(lib.root).as_posix() for p in lib.root.rglob("*") if "_Removed" in str(p))
    state_after = json.loads(lib.state_path.read_text())
    removed_flags = lambda st: {k for c in st["courses"].values() for k, it in c.get("items", {}).items() if it.get("removed")}  # noqa: E731
    check(before == after and removed_flags(state_before) == removed_flags(state_after),
          "nothing is retired when Canvas stops answering")
    check((new_week2 / "Duration Notes.pdf").exists() and (week3 / "Week 3 Overview.md").exists(), "library left intact")


# ============================================================== what the AI sees
LOCATOR_DOC = """---
title: Locator test
type: document
---

# Locator test

## [Slide 2] Immunization strategies

Match asset and liability sensitivities so that rate moves cancel out.

## Example

First worked example about coupons.

## Example

Second worked example about the liquidity premium.

## Chapter 1

### 1.1 Basics

Basics of bonds and their cash flows.

## Chapter 10

Advanced topics in term structure.

## [Lecture slides](https://example.com/slides.pdf)

Slides for the week are linked above.
"""


def ai_layer(fake) -> None:
    print("\n== What the AI sees: search, reading, study pass, packs, pictures")
    import asyncio

    from superstudent import notes
    from superstudent.compact import compact
    from superstudent.index import read_document, search, update_index
    from superstudent.library import Library
    from superstudent.mcp_server import _find_course, build_server
    from superstudent.outline import course_documents
    from superstudent.overview import EXAM_WORDS
    from superstudent.packs import make_pack
    from superstudent.render import RenderError, render

    lib = Library(WORK / "lib")
    fin = lib.root / "Fall 2026" / "FIN 6100 - Fixed Income"
    fin_rel = fin.relative_to(lib.root).as_posix()
    doc = fin / "My Files" / "locator-test.md"
    doc.write_text(LOCATOR_DOC, encoding="utf-8")
    rel = doc.relative_to(lib.root).as_posix()
    update_index(lib)

    hits = search(lib, "immunization strategies")
    check(any(h["path"] == rel and h["locator"].startswith("Slide 2") for h in hits),
          "slide titles are searchable", [(h["path"][-30:], h["locator"]) for h in hits])
    bare = [h for h in search(lib, "duration notes", limit=25) if h["snippet"].replace("**", "").strip() == h["title"]]
    check(not bare, "a document's title line is never a hit on its own", [h["path"][-40:] for h in bare])
    hit = next((h for h in search(lib, "liquidity premium") if h["path"] == rel), None)
    check(hit and hit["locator"] == "Example (2)" and "Second worked example" in read_document(lib, rel, locator=hit["locator"]),
          "a repeated heading's hit opens the right section", hit)
    section = read_document(lib, rel, locator="Chapter 1")
    check("Basics of bonds" in section and "Advanced topics" not in section, "'Chapter 1' opens chapter 1 (not 10), with its parts")
    hit = next((h for h in search(lib, "slides for the week") if h["path"] == rel), None)
    check(hit and hit["locator"] == "Lecture slides" and "linked above" in read_document(lib, rel, locator=hit["locator"]),
          "headings that are links open from their hit", hit)
    check("past the end" in read_document(lib, rel, start=10 ** 6), "reading past the end says so")
    paged = read_document(lib, rel, locator="Chapter 1", max_chars=30)
    check("Read on with start=30" in paged, "a long section is paged too")
    check(sum(1 for h in search(lib, "locator test") if h["path"] == rel) <= 1, "a title match counts once per document")
    quiz_hits = search(lib, "bond pricing", course="FIN", limit=8)
    check(sum(1 for h in quiz_hits if "Quiz 1" in h["path"]) <= 3, "one document's title doesn't crowd out the rest",
          [h["path"][-40:] for h in quiz_hits])
    check(_find_course(lib, "fin-6100") == fin_rel and _find_course(lib, "Fixed Income: Bonds") == "" and
          search(lib, "duration", course="FIN-6100"), "courses found by code or name, punctuation ignored")

    notes_text = "Slide 2: immunization matches sensitivities. The worked examples cover coupons and the liquidity premium. " * 3
    check(notes.save(lib, rel, notes_text, by="Claude")["ok"], "notes saved for the test document")
    status = notes.status_fn(lib)
    check(status(rel) == "current referenced notes", "current referenced notes are labeled accurately")
    doc.write_text(LOCATOR_DOC + "\n## Slide 3\n\nA new slide about key-rate durations.\n", encoding="utf-8")
    check(notes.status_fn(lib)(rel) == "changed since notes saved", "a changed source is flagged in the outline")
    r = notes.save(lib, rel, "Slide 2 " + "x" * 210000, by="Claude")
    check(not r["ok"] and "most one save can hold" in r["message"], "notes too long are refused, not cut short")

    docs = [d.rel for d in course_documents(lib, fin)]
    check(any(d.endswith("Quizzes/Quiz 1 Bond Pricing.md") for d in docs), "quizzes with past questions are in the study pass")
    check(not any("My Submissions" in d for d in docs), "the student's own submissions aren't")

    pack = Path(make_pack(lib, fin_rel, [2])["path"])
    mod = next(p for p in pack.iterdir() if p.name.startswith("02 - "))
    contents = read(mod / "_Module Contents.md")
    check((mod / "Linked" / "Quiz 1 Bond Pricing.md").exists() and "Linked/Quiz%201%20Bond%20Pricing.md" in contents,
          "exam packs include what a module links to elsewhere")

    server = build_server()

    def call(name: str, **args) -> str:
        res = asyncio.run(server.call_tool(name, args))
        blocks = res[0] if isinstance(res, tuple) else res
        blocks = getattr(blocks, "content", blocks)
        return "\n".join(getattr(b, "text", "") for b in blocks)

    out = call("course_file", course="FIN", file="outline:2")
    check("## Module 2" in out and "## Module 1" not in out, "one module of the outline at a time")
    check("Unknown course file" in call("course_file", course="FIN", file="bogus"), "an unknown course file is an error")
    check(call("search_course_materials", query="duration", course="no such course").startswith("No course matching"),
          "searching a course that doesn't exist says so")
    tools = asyncio.run(server.list_tools())
    check(all(getattr(t, "outputSchema", None) is None for t in tools if t.name != "view_page"),
          "results are plain text (no duplicate JSON copy)")

    screens = sorted((fin / "Modules").rglob("Week 2 Lecture Recording.mp4.assets/*.jpg"))
    check(bool(screens), "lecture screen snapshots exist")
    link = "Week%202%20Lecture%20Recording.mp4.assets/" + screens[0].name
    check(render(lib, link)[0].exists(), "a screen link copied from a transcript opens with view_page")
    from PIL import Image

    photo = fin / "My Files" / "phone_photo.jpg"
    shutil.copy2(WORK / "phone_photo.jpg", photo) if (WORK / "phone_photo.jpg").exists() else None
    if photo.exists():
        with Image.open(render(lib, photo.relative_to(lib.root).as_posix())[0]) as im:
            check(im.height > im.width, "a sideways phone photo is shown upright", im.size)
    for ext in (".bmp", ".tif"):
        pic = fin / "My Files" / f"chart{ext}"
        Image.new("RGB", (40, 30), "white").save(pic)
        check(render(lib, pic.relative_to(lib.root).as_posix())[0].suffix in (".png", ".jpg"), f"{ext} pictures can be viewed")
    (fin / "My Files" / "diagram.svg").write_text("<svg xmlns='http://www.w3.org/2000/svg'/>")
    try:
        render(lib, f"{fin_rel}/My Files/diagram.svg")
        refused = False
    except RenderError:
        refused = True
    check(refused, "SVG drawings get a clear message instead of a broken picture")
    pic_rel = f"{fin_rel}/My Files/chart.bmp"
    check(call("view_page", path=pic_rel, pages="1-3").count("Image files") <= 1 and
          len([b for b in (asyncio.run(server.call_tool("view_page", {"path": pic_rel, "pages": "1-3"})))
               if getattr(b, "type", "") == "image"]) <= 1, "a picture isn't sent three times")
    try:
        read_document(lib, screens[0].relative_to(lib.root).as_posix())
        binary = False
    except FileNotFoundError as exc:
        binary = "view_page" in str(exc)
    check(binary, "reading a picture as text points to view_page instead")
    lean = compact("---\ntitle: Scan\nnotes: no text layer; OCR found nothing\n---\n\n# Scan\n\n_No text found in this file._\n")
    check("no text layer" in lean, "a file with no text says why")

    outside = WORK / "outside.md"
    outside.write_text("# Secret\n\nnot course material\n", encoding="utf-8")
    link_md = fin / "My Files" / "linked.md"
    link_md.symlink_to(outside)
    update_index(lib)
    check(not search(lib, "secret"), "files linked from outside the library aren't indexed")
    try:
        read_document(lib, link_md.relative_to(lib.root).as_posix())
        blocked = False
    except FileNotFoundError:
        blocked = True
    check(blocked and "Current notes" in call("study_progress", course="FIN"), "…can't be read, and don't break the notes workflow")
    link_md.unlink()

    for text, want in (("The final value of the bond", False), ("the test statistic", False), ("comprehensive income", False),
                       ("The final will be cumulative.", True), ("Quiz 3 covers chapter 2", True)):
        check(bool(EXAM_WORDS.search(text)) == want, f"exam-hint wording: {text!r}")


# ============================================================== the app and installer
def app_layer() -> None:
    print("\n== The app: updates, transcription install, uninstall, schedule, files")
    import time as _time

    import superstudent.cli as cli
    import superstudent.gui.app as app
    import superstudent.scheduler as scheduler
    from superstudent.config import APP_DIR
    from superstudent.library import Library
    from superstudent.uninstall import stop_running_sync
    from superstudent.util import atomic_write_text

    APP_DIR.mkdir(parents=True, exist_ok=True)
    (APP_DIR / "just-updated").write_text("1.6.0\n")
    a = app.App(native=False)
    attention = a.state()["attention"]
    check(any(x.get("kind") == "updated" for x in attention), "after an update, the app says to reopen ChatGPT/Claude")
    a.dismiss_update()
    check(not (APP_DIR / "just-updated").exists(), "…until dismissed")

    old = app.App(native=False)
    old.serve(0)
    old.write_info()
    info = json.loads(app.INFO_FILE.read_text())
    info["version"] = "1.5.0"
    app.INFO_FILE.write_text(json.dumps(info))
    reused = app._focus_existing()
    _time.sleep(1.0)
    check(not reused and old.stopping, "a window left open from an older version is closed, not reused")
    old.shutdown()

    (APP_DIR / "bin").mkdir(parents=True, exist_ok=True)
    (APP_DIR / "bin" / "uv").write_text("#!/bin/sh\n")
    (APP_DIR / "constraints.txt").write_text("av==14.1.0\n")
    cmd = app.transcription_install_command()
    check(cmd[0].endswith("/bin/uv") and "--python" in cmd and "--constraint" in cmd, "transcription installs with uv", cmd)
    (APP_DIR / "bin" / "uv").unlink()
    (APP_DIR / "constraints.txt").unlink()

    lib = Library(WORK / "lib")
    holder = subprocess.Popen([sys.executable, "-c", "import sys, time; sys.path.insert(0, sys.argv[1]);"
                               "from superstudent.library import Library; from pathlib import Path\n"
                               "with Library(Path(sys.argv[2])).lock():\n    time.sleep(120)", str(ROOT), str(lib.root)])
    for _ in range(50):
        if lib.sync_running():
            break
        _time.sleep(0.1)
    check(lib.sync_running(), "a stand-in update is running")
    check(stop_running_sync(lib) and holder.wait(timeout=20) is not None and not lib.sync_running(),
          "uninstall stops a running update first")

    calls = []
    saved = (scheduler.platform.system, scheduler.subprocess.run, scheduler._plist_path, scheduler._sync_running)
    scheduler.platform.system = lambda: "Darwin"
    scheduler.subprocess.run = lambda args, **kw: (calls.append(list(args)), subprocess.CompletedProcess(args, 0, "", ""))[1]
    scheduler._plist_path = lambda: WORK / "agent.plist"
    scheduler._sync_running = lambda: True
    try:
        scheduler.schedule(6)                    # first time: loads it
        calls.clear()
        message = scheduler.schedule(12)         # changed while an update runs: postponed
        check(scheduler.PENDING.exists() and not any("bootout" in c for c in calls) and "starting when" in message,
              "changing the schedule doesn't stop a running update", message)
        scheduler._sync_running = lambda: False
        scheduler.apply_pending()
        check(not scheduler.PENDING.exists() and any("bootstrap" in c for c in calls), "…and applies once it's done")
        calls.clear()
        scheduler.schedule(12)
        check(not any("bootout" in c for c in calls), "an unchanged schedule leaves the job alone")
    finally:
        scheduler.platform.system, scheduler.subprocess.run, scheduler._plist_path, scheduler._sync_running = saved
        scheduler.PENDING.unlink(missing_ok=True)

    private = WORK / "config.toml"
    private.write_text("a = 1\n")
    os.chmod(private, 0o600)
    linked = WORK / "config-link.toml"
    linked.symlink_to(private)
    atomic_write_text(linked, "a = 2\n")
    check(linked.is_symlink() and private.read_text() == "a = 2\n" and (private.stat().st_mode & 0o777) == 0o600,
          "rewriting a private config keeps it private, and keeps a linked config linked")

    saved_mac = cli.IS_MAC
    cli.IS_MAC = True
    try:
        check(cli._tcc_warning(Path("/Volumes/USB/SuperStudent")) and
              cli._tcc_warning(Path.home() / "Library" / "CloudStorage" / "OneDrive" / "SuperStudent"),
              "external drives and cloud folders get the background-access warning")
    finally:
        cli.IS_MAC = saved_mac

    script = WORK / "lib" / "Fall 2026" / "FIN 6100 - Fixed Income" / "My Files" / "run me.command"
    script.write_text("#!/bin/sh\necho hi\n")
    a.actions.clear() if hasattr(a, "actions") else None
    a.open_item(script.relative_to(WORK / "lib").as_posix())
    check(a.actions and a.actions[-1]["open"][0] == "-R", "a runnable course file is shown in Finder, not run", a.actions[-1:])


def main() -> None:
    converters()
    from fake_canvas import FakeCanvas
    from make_fixtures import build_all

    fixtures = build_all(Path(os.environ.get("SS_TEST_DIR") or WORK) / "fixtures")
    fake = FakeCanvas(fixtures)
    base = fake.start()
    os.environ["SUPERSTUDENT_CANVAS_URL"] = base
    client(fake)
    keychain()
    syncing(fake, fixtures)
    ai_layer(fake)
    app_layer()
    fake.stop()
    print(f"\nALL {len(PASSED)} CHECKS PASSED" + (f" ({len(SKIPPED)} skipped)" if SKIPPED else "") + f"  (work dir: {WORK})")


if __name__ == "__main__":
    main()
