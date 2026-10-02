"""Build realistic course files for the fake Canvas: slides, PDFs, Word, Excel, an image, a short video."""

from __future__ import annotations

import io
from pathlib import Path


def make_pdf(path: Path) -> None:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Duration and Convexity", fontsize=20)
    page.insert_text((72, 110), "Macaulay duration is the weighted average time to receive cash flows.", fontsize=11)
    page.insert_text((72, 130), "Modified duration = Macaulay duration / (1 + y/k).", fontsize=11)
    page.insert_text((72, 150), "Exam tip: know how to compute modified duration by hand.", fontsize=11)
    page = doc.new_page()
    page.insert_text((72, 72), "Price sensitivity", fontsize=16)
    page.insert_text((72, 100), "dP/P = -D* x dy + 1/2 x C x dy^2", fontsize=11)
    page.insert_text((72, 120), "P = sum C/(1+y)^t", fontsize=11)
    page.insert_text((72, 140), "D = sum t w_t", fontsize=11)
    page.insert_text((72, 160), "C = sum t(t+1) w_t / (1+y)^2", fontsize=11)
    for i in range(40):  # a chart drawn with vector lines
        page.draw_line((100 + i * 8, 400), (100 + i * 8, 400 - (i * i) % 150))
    page = doc.new_page()  # image-only page (like a scan)
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 400, 300), 0)
    pix.clear_with(200)
    page.insert_image(pymupdf.Rect(72, 72, 500, 400), pixmap=pix)
    doc.save(str(path))


def make_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    layout = prs.slide_layouts[1]
    titles = ["Bond Basics", "Yield to Maturity", "Price-Yield Relationship"]
    for i, title in enumerate(titles, 1):
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text_frame.text = f"Key point {i}: prices and yields move inversely"
        slide.notes_slide.notes_text_frame.text = f"Speaker notes for {title}: emphasize this for the midterm."
    chart_data = CategoryChartData()
    chart_data.categories = ["1Y", "2Y", "5Y", "10Y"]
    chart_data.add_series("Yield %", (4.1, 3.9, 3.8, 4.2))
    prs.slides[2].shapes.add_chart(XL_CHART_TYPE.LINE, Inches(1), Inches(3.5), Inches(6), Inches(3), chart_data)
    img = io.BytesIO()
    from PIL import Image, ImageDraw

    im = Image.new("RGB", (800, 500), "white")
    d = ImageDraw.Draw(im)
    d.line([(50, 450), (750, 50)], fill="black", width=5)
    d.text((60, 60), "Yield curve", fill="black")
    im.save(img, format="PNG")
    img.seek(0)
    prs.slides[1].shapes.add_picture(img, Inches(1), Inches(3), Inches(4), Inches(2.5))
    # Move the last slide to the front: presentation order must win over internal file names.
    sld_ids = prs.slides._sldIdLst
    last = list(sld_ids)[-1]
    sld_ids.remove(last)
    sld_ids.insert(0, last)
    prs.save(str(path))


def make_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("Problem Set Handout", 1)
    doc.add_paragraph("Compute the price of a 5-year bond with a 6% coupon at a 5% yield.")
    doc.add_heading("Hints", 2)
    doc.add_paragraph("Use the annuity formula for coupons.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Year"
    table.cell(0, 1).text = "Cash flow"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "60"
    doc.save(str(path))


def make_xlsx(path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bond Model"
    ws.append(["Year", "Cash flow", "Discount factor", "PV"])
    for t in range(1, 6):
        cf = 60 if t < 5 else 1060
        ws.append([t, cf, f"=1/(1.05)^A{t + 1}", f"=B{t + 1}*C{t + 1}"])
    ws["F2"] = "Price"
    ws["G2"] = "=SUM(D2:D6)"
    wb.save(str(path))


def make_png(path: Path) -> None:
    from PIL import Image, ImageDraw

    im = Image.new("RGB", (640, 400), "white")
    d = ImageDraw.Draw(im)
    d.rectangle([40, 40, 600, 360], outline="black", width=4)
    d.text((60, 60), "Payoff diagram", fill="black")
    im.save(path)


def make_video(path: Path, seconds: int = 30) -> None:
    """A tiny 'lecture': three distinct slides, 10 s each, with a silent audio track."""
    import av
    import numpy as np
    from PIL import Image, ImageDraw

    container = av.open(str(path), mode="w")
    stream = container.add_stream("mpeg4", rate=2)
    stream.width, stream.height = 320, 240
    stream.pix_fmt = "yuv420p"
    stream.codec_context.gop_size = 4
    audio = container.add_stream("aac", rate=16000)
    colors = [(230, 230, 250), (250, 230, 200), (200, 240, 210)]
    for i in range(seconds * 2):
        slide = min(2, i // 20)
        im = Image.new("RGB", (320, 240), colors[slide])
        d = ImageDraw.Draw(im)
        d.rectangle([20 + slide * 40, 40, 140 + slide * 40, 200], fill=(40 * slide, 80, 160))
        d.text((30, 20), f"Slide {slide + 1}", fill="black")
        frame = av.VideoFrame.from_ndarray(np.asarray(im), format="rgb24")
        for packet in stream.encode(frame):
            container.mux(packet)
    for packet in stream.encode():
        container.mux(packet)
    samples = np.zeros((1, 1024), dtype=np.float32)
    try:
        layout = "mono"
        for n in range(int(seconds * 16000 / 1024)):
            af = av.AudioFrame.from_ndarray(samples, format="fltp", layout=layout)
            af.sample_rate = 16000
            af.pts = n * 1024
            for packet in audio.encode(af):
                container.mux(packet)
        for packet in audio.encode():
            container.mux(packet)
    except Exception:
        pass
    container.close()


def build_all(folder: Path) -> dict:
    folder.mkdir(parents=True, exist_ok=True)
    files = {
        "Lecture 1 - Bond Basics.pptx": make_pptx,
        "Duration Notes.pdf": make_pdf,
        "Problem Set Handout.docx": make_docx,
        "Bond Model.xlsx": make_xlsx,
        "payoff.png": make_png,
        "Week 2 Lecture Recording.mp4": make_video,
    }
    out = {}
    for name, fn in files.items():
        target = folder / name
        if not target.exists():
            fn(target)
        out[name] = target
    return out


def unlabeled_diagram(path: Path) -> Path:
    """A 'shoulder' drawing with no words in it (labels live on the slide as text boxes and arrows)."""
    from PIL import Image, ImageDraw

    im = Image.new("RGB", (900, 700), (252, 246, 240))
    d = ImageDraw.Draw(im)
    d.ellipse([300, 120, 620, 420], fill=(214, 120, 110), outline=(120, 40, 40), width=6)    # deltoid
    d.rectangle([420, 400, 500, 690], fill=(236, 226, 200), outline=(90, 80, 60), width=5)   # humerus
    d.polygon([(80, 80), (330, 150), (300, 260), (60, 220)], fill=(200, 90, 90), outline=(110, 30, 30))  # supraspinatus
    d.polygon([(60, 260), (300, 290), (320, 420), (90, 470)], fill=(190, 80, 80), outline=(110, 30, 30))  # infraspinatus
    im.save(path)
    return path


def make_anatomy_pptx(path: Path, work: Path) -> Path:
    """Slides like an anatomy lecture: a picture with text-box labels and arrows laid over it, a callout, a
    dark slide, a table, a grouped diagram and bullets."""
    from lxml import etree
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Emu, Inches, Pt

    def arrow(slide, x1, y1, x2, y2):
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.width = Pt(2)
        c.line.color.rgb = RGBColor(0, 0, 0)
        ln = c.line._get_or_add_ln()
        ln.append(etree.SubElement(ln, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", type="triangle"))
        return c

    def label(slide, text, x, y, w=2.2, h=0.5, size=20):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tb.text_frame.text = text
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
        tb.text_frame.paragraphs[0].runs[0].font.bold = True
        return tb

    pic = unlabeled_diagram(work / "shoulder-plain.png")
    prs = Presentation()
    blank = prs.slide_layouts[6]
    title_only = prs.slide_layouts[5]

    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "Posterior shoulder muscles"
    s.shapes.add_picture(str(pic), Inches(3.0), Inches(1.6), Inches(4.2), Inches(3.27))
    label(s, "Supraspinatus", 0.3, 1.6)
    label(s, "Infraspinatus", 0.3, 3.4)
    label(s, "Deltoid", 7.5, 1.9, 1.8)
    label(s, "Humerus", 7.5, 4.3, 1.8)
    arrow(s, 2.4, 1.85, 3.6, 2.2)
    arrow(s, 2.4, 3.65, 3.7, 3.4)
    arrow(s, 7.5, 2.15, 6.0, 2.4)
    arrow(s, 7.5, 4.55, 5.9, 4.6)
    call = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.2), Inches(5.9), Inches(3.8), Inches(0.9))
    call.fill.solid()
    call.fill.fore_color.rgb = RGBColor(255, 230, 80)
    call.text_frame.text = "Know all four rotator cuff muscles for the exam"
    call.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
    call.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 0, 0)
    s.notes_slide.notes_text_frame.text = "SITS muscles: supraspinatus, infraspinatus, teres minor, subscapularis."

    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor(20, 30, 60)
    tb = s.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(2))
    tb.text_frame.text = "Rotator cuff tears are the most common shoulder injury"
    run = tb.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(36)
    run.font.color.rgb = RGBColor(255, 255, 255)

    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "Muscle actions"
    rows = [("Muscle", "Action", "Nerve"), ("Supraspinatus", "Abduction (first 15°)", "Suprascapular"),
            ("Infraspinatus", "External rotation", "Suprascapular"), ("Deltoid", "Abduction", "Axillary")]
    table = s.shapes.add_table(len(rows), 3, Inches(0.7), Inches(1.8), Inches(8.6), Inches(2.4)).table
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            table.cell(r, c).text = val

    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "Nerve pathway"
    boxes = []
    for i, name in enumerate(["C5 root", "Upper trunk", "Suprascapular nerve"]):
        b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6 + i * 3.1), Inches(3), Inches(2.5), Inches(1))
        b.text_frame.text = name
        boxes.append(b)
    group = s.shapes.add_group_shape(boxes)
    for i in range(2):
        arrow(s, 3.1 + i * 3.1, 3.5, 3.7 + i * 3.1, 3.5)

    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Summary"
    body = s.placeholders[1].text_frame
    body.text = "Rotator cuff = four muscles (SITS)"
    for line, level in (("Stabilize the humeral head in the glenoid", 1), ("Deltoid is not part of the cuff", 0)):
        para = body.add_paragraph()
        para.text = line
        para.level = level
    prs.save(str(path))
    return path
