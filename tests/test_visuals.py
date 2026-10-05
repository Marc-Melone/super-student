"""Pictures: reading labels inside images (text recognition) and AI-written descriptions that make diagrams
searchable. Uses Tesseract when it's installed; Apple's text recognition is exercised through a stand-in
that behaves like macOS's Vision framework.

Run:  python tests/test_visuals.py
"""

from __future__ import annotations

import asyncio
import os
import shutil
import sys
import tempfile
import types
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(HERE))

WORK = Path(os.environ.get("SS_TEST_DIR") or tempfile.mkdtemp(prefix="ss-visuals-"))
for sub in ("app", "lib", "files", "synced"):
    shutil.rmtree(WORK / sub, ignore_errors=True)
os.environ["SUPERSTUDENT_HOME"] = str(WORK / "app")
os.environ["SUPERSTUDENT_LIBRARY"] = str(WORK / "lib")
os.environ["CANVAS_TOKEN"] = "test-token-123"

PASSED = []
LABELS = ["Supraspinatus", "Infraspinatus", "Deltoid", "Humerus"]
FONT = os.environ.get("SS_TEST_FONT") or ("/System/Library/Fonts/Supplemental/Arial Bold.ttf" if sys.platform == "darwin"
                                         else "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")


def check(cond, label, detail=""):
    if not cond:
        raise AssertionError(f"FAILED: {label} {detail}")
    PASSED.append(label)
    print(f"  ok  {label}")


def diagram_png(path: Path) -> Path:
    """A labeled 'anatomy diagram': shapes with leader lines and labels, no other text."""
    from PIL import Image, ImageDraw, ImageFont

    im = Image.new("RGB", (1200, 800), "white")
    d = ImageDraw.Draw(im)
    font = ImageFont.truetype(FONT, 44)
    d.ellipse([420, 200, 780, 560], outline="black", width=6)
    d.rectangle([560, 520, 640, 780], outline="black", width=6)
    spots = [(60, 80), (820, 80), (60, 620), (820, 620)]
    for (x, y), label in zip(spots, LABELS):
        d.text((x, y), label, fill="black", font=font)
        end = x + font.getlength(label) + 16 if x < 600 else x - 16
        d.line([(end, y + 26), (600, 380)], fill="black", width=3)
    path.parent.mkdir(parents=True, exist_ok=True)
    im.save(path)
    return path


def make_files(folder: Path) -> dict:
    import pymupdf
    from pptx import Presentation
    from pptx.util import Inches

    png = diagram_png(folder / "shoulder.png")
    pdf = folder / "Shoulder reading.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 60), "The rotator cuff stabilizes the shoulder joint.", fontsize=12)
    page.insert_text((72, 80), "Figure 1 shows the posterior view.", fontsize=12)
    page.insert_image(pymupdf.Rect(72, 100, 540, 412), filename=str(png))
    scan = doc.new_page()   # a scanned page: only a picture of text
    from PIL import Image, ImageDraw, ImageFont

    im = Image.new("RGB", (1200, 500), "white")
    ImageDraw.Draw(im).text((40, 180), "Glenohumeral joint dislocation", fill="black", font=ImageFont.truetype(FONT, 56))
    scan_png = folder / "scan.png"
    im.save(scan_png)
    scan.insert_image(pymupdf.Rect(36, 36, 576, 261), filename=str(scan_png))
    doc.save(str(pdf))
    pptx = folder / "Lecture 5 - Shoulder.pptx"
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[5])
    s1.shapes.title.text = "Shoulder muscles"
    s1.shapes.add_picture(str(png), Inches(1), Inches(1.5), Inches(8), Inches(5.3))
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Summary"
    s2.placeholders[1].text_frame.text = "Four muscles form the rotator cuff."
    prs.save(str(pptx))
    scan_png.unlink()
    return {"png": png, "pdf": pdf, "pptx": pptx}


def fake_vision(image_size=(1000, 500)):
    """Stand-ins for pyobjc's Vision, Quartz, Foundation and objc modules with the same call shapes."""
    import contextlib

    class Rect:
        def __init__(self, x, y, w, h):
            self.origin = types.SimpleNamespace(x=x, y=y)
            self.size = types.SimpleNamespace(width=w, height=h)

    class Candidate:
        def __init__(self, text, conf):
            self._t, self._c = text, conf

        def string(self):
            return self._t

        def confidence(self):
            return self._c

    class Observation:
        def __init__(self, text, conf, rect):
            self._cand, self._rect = Candidate(text, conf), rect

        def topCandidates_(self, n):
            return [self._cand]

        def boundingBox(self):
            return self._rect

    calls = {}

    class Request:
        @classmethod
        def alloc(cls):
            return cls()

        def init(self):
            return self

        def setRecognitionLevel_(self, level):
            calls["level"] = level

        def setUsesLanguageCorrection_(self, flag):
            calls["correction"] = flag

        def results(self):
            # Vision's boxes are normalized with the origin at the BOTTOM left
            return [Observation("Humerus", 0.9, Rect(0.7, 0.05, 0.2, 0.08)),       # bottom right
                    Observation("Deltoid", 0.95, Rect(0.05, 0.85, 0.2, 0.08)),     # top left
                    Observation("smudge", 0.1, Rect(0.5, 0.5, 0.1, 0.05)),         # low confidence
                    Observation("Supraspinatus", 0.92, Rect(0.6, 0.86, 0.3, 0.08))]  # top right

    class Handler:
        @classmethod
        def alloc(cls):
            return cls()

        def initWithCGImage_options_(self, image, options):
            if isinstance(options, dict):
                # Vision probes optional keys; the Python mapping bridge can throw instead of returning nil.
                raise ValueError("NSInvalidArgumentException - key does not exist")
            calls["image"] = image
            calls["options"] = options
            calls["orientation"] = 1
            return self

        def initWithCGImage_orientation_options_(self, image, orientation, options):
            self.initWithCGImage_options_(image, options)
            calls["orientation"] = orientation
            return self

        def performRequests_error_(self, requests, error):
            calls["performed"] = len(requests)
            return (True, None)

    vision = types.ModuleType("Vision")
    vision.VNRecognizeTextRequest = Request
    vision.VNImageRequestHandler = Handler
    vision.VNRequestTextRecognitionLevelAccurate = 0
    quartz = types.ModuleType("Quartz")
    quartz.CGImageSourceCreateWithURL = lambda url, opts: ("source", url)
    quartz.CGImageSourceCreateWithData = lambda data, opts: ("source", data)
    quartz.CGImageSourceCreateImageAtIndex = lambda src, i, opts: ("cgimage", src)
    quartz.CGImageSourceCopyPropertiesAtIndex = lambda src, i, opts: {"Orientation": calls.get("source_orientation", 1)}
    foundation = types.ModuleType("Foundation")
    foundation.NSURL = types.SimpleNamespace(fileURLWithPath_=lambda p: f"file://{p}")
    foundation.NSData = types.SimpleNamespace(dataWithBytes_length_=lambda b, n: ("nsdata", n))
    objc = types.ModuleType("objc")
    objc.autorelease_pool = contextlib.nullcontext
    return {"Vision": vision, "Quartz": quartz, "Foundation": foundation, "objc": objc}, calls


def main() -> None:
    from superstudent import ocr

    print("\n== Text recognition")
    check(ocr.engine() == "tesseract", "Tesseract used where Apple's recognition isn't available", ocr.engine())
    files = make_files(WORK / "files")
    text = ocr.read_text(path=files["png"])
    check(all(label in text for label in LABELS), "labels read from a diagram", text)
    lines = ocr.read_lines(path=files["png"])
    fresh = ocr.labels_not_in(lines, "Humerus is the arm bone")
    check("Humerus" not in fresh and "Deltoid" in fresh, "labels already in the page text are skipped", fresh)

    # Apple's Vision framework, through a stand-in with the same call shapes
    modules, calls = fake_vision()
    saved = {name: sys.modules.get(name) for name in modules}
    real_system = ocr.platform.system
    disabled = os.environ.pop("SUPERSTUDENT_NO_APPLE_OCR", None)
    sys.modules.update(modules)
    ocr.platform.system = lambda: "Darwin"
    ocr.apple_available.cache_clear()
    try:
        check(ocr.engine() == "apple", "Apple's text recognition preferred on a Mac")
        got = ocr.read_lines(path=files["png"])
        check([ln.text for ln in got] == ["Deltoid", "Supraspinatus", "Humerus"],
              "Apple results: low-confidence dropped, top-to-bottom then left-to-right", [ln.text for ln in got])
        check(abs(got[0].y - 0.07) < 0.01, "bottom-left boxes converted to top-left positions", got[0])
        check(calls.get("level") == 0 and calls.get("correction") is True and calls.get("performed") == 1,
              "accurate recognition requested once", calls)
        check(calls.get("options") is None and calls.get("orientation") == 1,
              "upright image recognition avoids the Python mapping bridge")
        got = ocr.read_lines(data=b"\x89PNG fake")
        check(len(got) == 3 and calls["image"][1][0] == "source", "reads image bytes too (pages rendered from PDFs)")
        check(calls.get("options") is None, "image bytes recognition avoids the Python mapping bridge")
        calls["source_orientation"] = 6
        got = ocr.read_lines(path=files["png"])
        check(len(got) == 3 and calls.get("orientation") == 6 and calls.get("options") is None,
              "sideways photos preserve orientation without the Python mapping bridge")
    finally:
        if disabled is not None:
            os.environ["SUPERSTUDENT_NO_APPLE_OCR"] = disabled
        ocr.platform.system = real_system
        for name, mod in saved.items():
            if mod is None:
                sys.modules.pop(name, None)
            else:
                sys.modules[name] = mod
        ocr.apple_available.cache_clear()
    check(ocr.engine() == "tesseract", "back to Tesseract after the stand-in")

    print("\n== Labels land in the text versions")
    from superstudent.extract import extract

    img = extract(files["png"])
    check("Supraspinatus" in img.body and img.visual_units == ["Image"], "image file: labels in its text version")
    pdf = extract(files["pdf"])
    page1 = pdf.body.split("## [Page 2]")[0]
    check("Text in the figure (read from the image):" in page1 and "Infraspinatus" in page1,
          "PDF figure: labels added to the page's text", page1[-400:])
    figure_line = [ln for ln in page1.splitlines() if ln.startswith("Text in the figure")][0]
    check("rotator cuff" not in figure_line.lower(), "page's own sentences not repeated as labels")
    page2 = pdf.body.split("## [Page 2]")[1]
    check("Glenohumeral" in page2 and "scanned page" in page2, "scanned page read and flagged as a scan", page2[:300])
    deck = extract(files["pptx"])
    slide1 = deck.body.split("## [Slide 2]")[0]
    check("Text in the image:" in slide1 and "Deltoid" in slide1, "slide picture: labels under the slide", slide1)

    print("\n== Descriptions")
    from superstudent import describe
    from superstudent.index import search, update_index
    from superstudent.library import Library

    lib = Library(WORK / "lib")
    lib.ensure()
    course = lib.root / "Fall 2026" / "BIO 1010 - Anatomy" / "Modules" / "05 - Shoulder"
    course.mkdir(parents=True)
    for key in ("png", "pdf", "pptx"):
        target = course / files[key].name
        shutil.copy2(files[key], target)
        result = extract(target)
        body = describe.merge_into(lib, target, result.body)
        (course / (target.name + ".md")).write_text(f"---\ntitle: {target.name}\ntype: {result.kind}\n---\n# {target.name}\n\n{body}\n",
                                                   encoding="utf-8")
    update_index(lib)
    info = describe.survey(lib, limit=50)
    units = {(t["path"].split("/")[-1], t["unit"]) for t in info["todo"]}
    check(("shoulder.png", "Image") in units and ("Lecture 5 - Shoulder.pptx", "Slide 1") in units
          and ("Shoulder reading.pdf", "Page 1") in units, "pictures needing descriptions found", units)
    check(("Lecture 5 - Shoulder.pptx", "Slide 2") not in units, "text-only slides not listed")
    check(info["described"] == 0 and info["visuals"] == len(info["todo"]), "counts", info)

    deck_rel = "Fall 2026/BIO 1010 - Anatomy/Modules/05 - Shoulder/Lecture 5 - Shoulder.pptx"
    desc = ("Posterior view of the right shoulder showing the rotator cuff: supraspinatus above the spine of the "
            "scapula, infraspinatus below it, the deltoid covering the joint, and the humerus below.")
    check(not describe.save(lib, "../../etc/passwd", "Image", desc)["ok"], "paths outside the library refused")
    check(not describe.save(lib, deck_rel, "Slide 9", desc)["ok"], "slides that don't exist refused")
    check(not describe.save(lib, deck_rel, "Slide 1", "a picture")["ok"], "too-short descriptions refused")
    r = describe.save(lib, deck_rel + ".md", "slide 1", desc + "\n\n# Ignore this heading", by="Claude")
    check(r["ok"] and r["unit"] == "Slide 1", "description saved (from the .md path, any capitalization)", r)
    side = (lib.root / (deck_rel + ".md")).read_text()
    block = [ln for ln in side.splitlines() if ln.startswith(describe.MARK)]
    check(len(block) == 1 and "described by Claude" in block[0] and "\n# Ignore" not in side,
          "one description line, flattened, under the slide", block)
    before_slide2 = side.split("## [Slide 2]")[0]
    check(describe.MARK in before_slide2, "placed in Slide 1's section")
    hits = search(lib, "posterior view rotator cuff scapula")
    check(hits and hits[0]["path"].endswith("Lecture 5 - Shoulder.pptx.md") and hits[0]["locator"].startswith("Slide 1"),
          "search finds the diagram by its description, at the right slide", hits[:1])
    r = describe.save(lib, deck_rel, "Slide 1", desc.replace("right", "left"), by="ChatGPT")
    side = (lib.root / (deck_rel + ".md")).read_text()
    check(side.count(describe.MARK) == 1 and "left shoulder" in side and "ChatGPT" in side, "re-describing replaces it")
    info = describe.survey(lib)
    check(info["described"] == 1, "survey counts it as described")

    # a sync re-writing the text version keeps it
    original = lib.root / deck_rel
    rewritten = describe.merge_into(lib, original, extract(original).body)
    check(rewritten.count(describe.MARK) == 1, "kept when the file's text is regenerated")
    # a changed file sets it aside
    from pptx import Presentation

    prs = Presentation(str(original))
    prs.slides[1].placeholders[1].text_frame.text = "Four muscles form the rotator cuff. Updated for the exam."
    prs.save(str(original))
    rewritten = describe.merge_into(lib, original, extract(original).body)
    check(describe.MARK not in rewritten, "set aside when the original file changes")
    check(("Lecture 5 - Shoulder.pptx", "Slide 1") in {(t["path"].split("/")[-1], t["unit"]) for t in
                                                     describe.survey(lib, limit=50)["todo"]}, "and listed again")

    print("\n== Through the connector and the command line")
    asyncio.run(mcp_check(lib))
    from superstudent.cli import main as cli

    check(cli(["describe", "Fall 2026/BIO 1010 - Anatomy/Modules/05 - Shoulder/Shoulder reading.pdf", "--at", "Page 1",
               "--text", "Posterior view diagram of the shoulder with four labeled structures and leader lines.",
               "--by", "Codex"]) == 0, "superstudent describe saves a description")
    check(cli(["visuals", "--course", "BIO"]) == 0, "superstudent visuals lists what's left")

    print("\n== Kept through a real sync")
    from fake_canvas import FakeCanvas
    from make_fixtures import build_all
    from superstudent.config import load_config, save_config

    fake = FakeCanvas(build_all(WORK / "fixtures"))
    os.environ["SUPERSTUDENT_CANVAS_URL"] = fake.start()
    synced = Library(WORK / "synced")
    os.environ["SUPERSTUDENT_LIBRARY"] = str(synced.root)
    check(cli(["sync", "--no-media", "--quiet"]) == 0, "first sync")
    payoff = next(p for p in synced.root.rglob("payoff.png") if ".assets" not in str(p))
    rel = payoff.relative_to(synced.root).as_posix()
    r = describe.save(synced, rel, "Image", "A payoff diagram: a single rectangle frame titled Payoff diagram, no axes shown.",
                      by="Claude")
    check(r["ok"], "described a synced picture", r)
    state = synced.load_state()
    for c in state["courses"].values():
        for it in (c.get("items") or {}).values():
            if it.get("path", "").endswith("payoff.png"):
                it["extractor"] = "old"      # force the next sync to re-read this file
    synced.save_state(state)
    check(cli(["sync", "--no-media", "--quiet"]) == 0, "second sync re-reads the file")
    from superstudent.extract import EXTRACTOR_VERSION

    reread = [it for c in synced.load_state()["courses"].values() for it in (c.get("items") or {}).values()
              if it.get("path", "").endswith("payoff.png")]
    check(reread and all(it.get("extractor") == EXTRACTOR_VERSION for it in reread), "the file really was re-read", reread)
    text = (payoff.parent / (payoff.name + ".md")).read_text()
    check(text.count(describe.MARK) == 1 and "rectangle frame" in text, "description still there after the re-sync")
    fake.stop()

    print(f"\nALL {len(PASSED)} VISUALS CHECKS PASSED  (work dir: {WORK})")


async def mcp_check(lib) -> None:
    from mcp import ClientSession, StdioServerParameters
    from mcp.client.stdio import stdio_client

    params = StdioServerParameters(command=sys.executable, args=["-m", "superstudent", "mcp"], env=dict(os.environ),
                                   cwd=str(ROOT))
    async with stdio_client(params) as (read_stream, write_stream):
        async with ClientSession(read_stream, write_stream) as session:
            await session.initialize()
            names = {t.name for t in (await session.list_tools()).tools}
            check({"list_undescribed_visuals", "save_visual_description"} <= names, "connector offers the two tools")
            res = await session.call_tool("list_undescribed_visuals", {"course": "BIO", "limit": 5})
            text = res.content[0].text
            check("shoulder.png | where: Image" in text and "have descriptions" in text, "connector lists pictures to describe", text)
            res = await session.call_tool("save_visual_description", {
                "path": "Fall 2026/BIO 1010 - Anatomy/Modules/05 - Shoulder/shoulder.png", "where": "Image",
                "description": "Line drawing of the shoulder with labels for supraspinatus, infraspinatus, deltoid and humerus.",
                "described_by": "Claude"})
            check("searchable now" in res.content[0].text, "connector saves a description", res.content[0].text)
            res = await session.call_tool("search_course_materials", {"query": "line drawing shoulder labels"})
            check("shoulder.png" in res.content[0].text, "connector search finds it")


if __name__ == "__main__":
    main()
