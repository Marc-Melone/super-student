"""Course files -> Markdown an AI can search, with page/slide markers and saved visuals.

Every extracted document uses the same locator headings, so search results can cite
exactly where something came from:

    ## [Page 12]            PDFs, e-books, scanned handouts
    ## [Slide 7] Title      PowerPoint decks (in presentation order, with speaker notes)
    ## [Sheet: Model]       Spreadsheets (cells keep their A1 addresses, formulas listed)
    ## [00:32:10]           Lecture transcripts

Pages or slides whose meaning lives in a picture (charts, diagrams, equations, scans)
are flagged with a "> Visual content:" line so the AI knows to look at the page image.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
import shutil
import statistics
import subprocess
import tempfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional, Set

from .content import simple_html_to_markdown
from .util import md_link

from . import ocr as _ocr
from .layout import column_order as _column_order

# Bumped when extraction improves; files are re-read from the copies already on disk (no re-download).
# The text-recognition engine is part of it, so gaining one (e.g. the Mac app) re-reads pictures once.
EXTRACTOR_VERSION = "4" + (f"+{_ocr.engine()}" if _ocr.engine() else "")   # 4: columns, equations, small pictures

MEDIA_EXT = {".mp4", ".mov", ".m4v", ".mkv", ".webm", ".avi", ".wmv", ".flv", ".mpeg", ".mpg", ".3gp",
             ".mp3", ".m4a", ".wav", ".aac", ".ogg", ".oga", ".opus", ".flac", ".wma", ".aiff", ".aif"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff", ".heic", ".heif", ".svg"}
VIEWABLE_IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp"}
PDFLIKE_EXT = {".pdf", ".epub", ".xps", ".oxps", ".mobi", ".fb2", ".cbz"}
SHEET_EXT = {".xlsx", ".xlsm", ".xltx", ".xltm"}
LEGACY_EXT = {".doc", ".rtf", ".odt", ".ppt", ".pps", ".ppsx", ".odp", ".xls", ".ods", ".key", ".pages", ".numbers", ".wpd"}
CODE_LANG = {
    ".py": "python", ".r": "r", ".rmd": "markdown", ".m": "matlab", ".sql": "sql", ".sas": "sas", ".do": "stata",
    ".jl": "julia", ".js": "javascript", ".ts": "typescript", ".java": "java", ".c": "c", ".cpp": "cpp",
    ".h": "c", ".cs": "csharp", ".go": "go", ".rs": "rust", ".rb": "ruby", ".sh": "bash", ".scala": "scala",
    ".tex": "latex", ".bib": "bibtex", ".json": "json", ".xml": "xml", ".yaml": "yaml", ".yml": "yaml",
    ".css": "css", ".toml": "toml", ".ini": "ini", ".cfg": "ini", ".vba": "vb", ".bas": "vb",
}
PLAIN_EXT = {".txt", ".md", ".markdown", ".log", ".text"}
MATH_CHARS = set("∑∫∂√±×÷≤≥≠≈∞∝∆∇∈∉⊂⊆∪∩→←⇒⇔′″αβγδεζηθικλμνξπρστυφχψωΓΔΘΛΞΠΣΦΨΩ")
LIGATURES = {"ﬁ": "fi", "ﬂ": "fl", "ﬀ": "ff", "ﬃ": "ffi", "ﬄ": "ffl", "\u00ad": "", "\u200b": ""}

MAX_TEXT_BYTES = 2 * 1024 * 1024
MAX_SHEET_ROWS = 300
MAX_SHEET_COLS = 60


@dataclass
class Extracted:
    kind: str                      # slides | pdf | document | spreadsheet | image | text | notebook | web | archive | unsupported
    body: str
    units: str = ""
    notes: List[str] = field(default_factory=list)
    visual_units: List[str] = field(default_factory=list)
    assets: List[str] = field(default_factory=list)


def kind_of(path: Path, content_type: Optional[str] = None) -> str:
    ext = path.suffix.lower()
    ctype = (content_type or "").lower()
    if ext in MEDIA_EXT or ctype.startswith(("video/", "audio/")):
        return "media"
    if ext in IMAGE_EXT or ctype.startswith("image/"):
        return "image"
    if ext in PDFLIKE_EXT or ctype == "application/pdf":
        return "pdf"
    if ext == ".pptx" or ext == ".potx":
        return "slides"
    if ext in (".docx", ".dotx", ".docm"):
        return "document"
    if ext in SHEET_EXT:
        return "spreadsheet"
    if ext in (".csv", ".tsv"):
        return "csv"
    if ext == ".ipynb":
        return "notebook"
    if ext in (".html", ".htm", ".xhtml"):
        return "web"
    if ext == ".zip":
        return "archive"
    if ext in LEGACY_EXT:
        return "legacy"
    if ext in CODE_LANG or ext in PLAIN_EXT or ctype.startswith("text/"):
        return "text"
    return "unknown"


def assets_dir_for(path: Path) -> Path:
    return path.parent / (path.name + ".assets")


def extract(path: Path, content_type: Optional[str] = None) -> Extracted:
    kind = kind_of(path, content_type)
    handlers = {
        "pdf": _pdf, "slides": _pptx, "document": _docx, "spreadsheet": _xlsx, "csv": _csv,
        "notebook": _notebook, "web": _html_file, "text": _text, "image": _image, "legacy": _legacy,
    }
    if kind in handlers:
        return handlers[kind](path)
    if kind == "archive":
        return Extracted("archive", "", notes=["Zip archive; its contents are unpacked next to it."])
    if kind == "media":
        return Extracted("media", "", notes=["Audio/video: see the transcript."])
    # Unknown extension: sniff for PDF or text.
    head = path.read_bytes()[:2048] if path.exists() else b""
    if head.startswith(b"%PDF"):
        return _pdf(path)
    if head and b"\x00" not in head:
        return _text(path)
    return Extracted("unsupported", "", notes=[f"No text extractor for '{path.suffix or 'this'}' files. Open the original."])


# ---------------------------------------------------------------- helpers

def _clean_text(text: str) -> str:
    for bad, good in LIGATURES.items():
        text = text.replace(bad, good)
    text = re.sub(r"(\w)-\n([a-z])", r"\1\2", text)          # re-join hyphenated line breaks
    lines = [ln.rstrip() for ln in text.splitlines()]
    out: List[str] = []
    blank = 0
    for ln in lines:
        if not ln.strip():
            blank += 1
            if blank <= 1:
                out.append("")
            continue
        blank = 0
        if ln.lstrip().startswith("#"):
            ln = ln.replace("#", "\\#", 1)                      # keep our locator headings unambiguous
        out.append(ln)
    return "\n".join(out).strip()


def _mathy(text: str) -> bool:
    if not text:
        return False
    symbols = sum(1 for ch in text if ch in MATH_CHARS)
    eq_lines = sum(1 for ln in text.splitlines() if "=" in ln and 3 <= len(ln.strip()) <= 90)
    return symbols >= 3 or eq_lines >= 3


def tesseract_available() -> bool:
    """True when some text recognition is available (Apple's on a Mac, or Tesseract). Name kept for callers."""
    return _ocr.available()


def ocr_image_file(path: Path, timeout: int = 90) -> str:
    return _clean_text(_ocr.read_text(path=Path(path), timeout=timeout)) if _ocr.available() else ""


def _ocr_budget(apple: int, tesseract: int) -> int:
    return {"apple": apple, "tesseract": tesseract}.get(_ocr.engine(), 0)


def _safe_inline(text: str) -> str:
    """OCR text going into a Markdown line: no brackets or heading marks that would change its meaning."""
    return re.sub(r"[\[\]()<>#*_`|]", " ", text).replace("  ", " ").strip()


def soffice_path() -> Optional[str]:
    if os.environ.get("SUPERSTUDENT_NO_SOFFICE"):     # tests: behave as if LibreOffice isn't installed
        return None
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def convert_with_soffice(path: Path, target_ext: str, out_dir: Path, timeout: int = 240,
                         convert_to: Optional[str] = None) -> Optional[Path]:
    """Convert with LibreOffice if it's installed (e.g. .ppt -> .pptx, .pptx -> .pdf)."""
    exe = soffice_path()
    if not exe:
        return None
    out_dir.mkdir(parents=True, exist_ok=True)
    profile = Path(tempfile.mkdtemp(prefix="ss-lo-"))
    try:
        subprocess.run(
            [exe, f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to",
             convert_to or target_ext.lstrip("."), "--outdir", str(out_dir), str(path)],
            capture_output=True, timeout=timeout,
        )
    except (OSError, subprocess.SubprocessError):
        return None
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    candidate = out_dir / (path.stem + "." + target_ext.lstrip("."))
    return candidate if candidate.exists() else None


PRESENTATION_EXT = {".pptx", ".ppt", ".pps", ".ppsx", ".odp", ".key"}
# LibreOffice leaves hidden slides out of PDFs unless told otherwise; the text version numbers them, so page N
# of the picture must be slide N.
WITH_HIDDEN_SLIDES = 'pdf:impress_pdf_Export:{"ExportHiddenSlides":{"type":"boolean","value":"true"}}'


def office_pdf(path: Path, out_dir: Path) -> Optional[Path]:
    """A PDF of an Office file made by LibreOffice, with every slide of a deck (hidden ones too) in order.
    None if LibreOffice isn't there, or if a deck's PDF doesn't have one page per slide (older LibreOffice
    versions ignore the hidden-slides option); the caller then draws the slides itself."""
    if path.suffix.lower() not in PRESENTATION_EXT:
        return convert_with_soffice(path, ".pdf", out_dir)
    pdf = convert_with_soffice(path, ".pdf", out_dir, convert_to=WITH_HIDDEN_SLIDES)
    if pdf is None or path.suffix.lower() not in (".pptx", ".ppsx"):
        return pdf
    try:
        import pymupdf
        from pptx import Presentation

        slides = len(Presentation(str(path)).slides)
        with pymupdf.open(str(pdf)) as doc:
            pages = doc.page_count
    except Exception:
        return pdf
    if pages != slides:
        try:
            pdf.unlink()
        except OSError:
            pass
        return None
    return pdf


def _md_table(rows: List[List[str]], header: Optional[List[str]] = None) -> str:
    if not rows and not header:
        return ""
    width = max([len(r) for r in rows] + [len(header or [])])

    def cell(v: object) -> str:
        return str(v if v is not None else "").replace("|", "\\|").replace("\n", " ").strip()

    head = [cell(h) for h in (header or rows[0])] + [""] * width
    head = head[:width]
    body = rows if header else rows[1:]
    lines = ["| " + " | ".join(h or " " for h in head) + " |", "|" + "---|" * width]
    for r in body:
        vals = [cell(v) for v in r] + [""] * width
        lines.append("| " + " | ".join(vals[:width]) + " |")
    return "\n".join(lines)


def _clear_assets(assets: Path) -> None:
    # A directory link must be replaced, not left behind after rmtree fails:
    # its children can themselves link outside the library.
    if assets.is_symlink():
        assets.unlink()
    elif assets.exists():
        shutil.rmtree(assets)


def _save_asset(assets: Path, name: str, data: bytes) -> str:
    assets.mkdir(parents=True, exist_ok=True)
    target = assets / name
    target.write_bytes(data)
    return f"{assets.name}/{name}"


def _to_png(data: bytes) -> Optional[bytes]:
    try:
        from PIL import Image
    except ImportError:
        return None
    try:
        with Image.open(io.BytesIO(data)) as img:
            buf = io.BytesIO()
            img.convert("RGB").save(buf, format="PNG")
            return buf.getvalue()
    except Exception:
        return None


# ---------------------------------------------------------------- PDF (and e-books)

def _page_text(page) -> str:
    """A PDF page's text in reading order. Columns of running text are kept apart (sorting every line by
    height would interleave a two-column article line by line), and cells of a table row stay on one line."""
    try:
        raw = page.get_text("dict")["blocks"]
    except Exception:
        return page.get_text("text") or ""
    blocks = []
    for block in raw:
        if block.get("type") != 0:
            continue
        rows: List[list] = []        # [y0, x1, text, size]
        for line in block.get("lines") or []:
            text = "".join(span.get("text", "") for span in line.get("spans") or [])
            if not text.strip():
                continue
            x0, y0, x1, _ = line.get("bbox", (0, 0, 0, 0))
            size = max((span.get("size", 10) for span in line.get("spans") or []), default=10)
            if rows and abs(y0 - rows[-1][0]) < 0.5 * size and x0 > rows[-1][1]:
                gap = x0 - rows[-1][1]
                rows[-1][2] += (" | " if gap > 1.5 * size else " ") + text.strip()
                rows[-1][1] = x1
            else:
                rows.append([y0, x1, text.strip(), size])
        if rows:
            x0, y0, x1, y1 = block.get("bbox", (0, 0, 0, 0))
            blocks.append((x0, y0, x1, y1, "\n".join(r[2] for r in rows)))
    if not blocks:
        return ""
    left = min(b[0] for b in blocks)
    width = max(1.0, max(b[2] for b in blocks) - left)
    ordered = _column_order(blocks, lambda b: b[2] - b[0], lambda b: b[4], width)
    return "\n".join(b[4] for b in ordered)


def _pdf(path: Path) -> Extracted:
    import pymupdf

    try:
        doc = pymupdf.open(str(path))
    except Exception as exc:
        return Extracted("pdf", "", notes=[f"Couldn't open this PDF ({exc.__class__.__name__}). Open the original."])
    with doc:
        if doc.needs_pass:
            return Extracted("pdf", "", notes=["Password-protected PDF; open the original."])
        count = doc.page_count
        long_doc = count > 200
        # Images that repeat on most pages are template decoration (logos, backgrounds): ignore them.
        xref_pages: Counter = Counter()
        drawing_counts: List[int] = []
        for page in doc:
            try:
                xref_pages.update({img[0] for img in page.get_images(full=False)})
            except Exception:
                pass
            if not long_doc:
                try:
                    drawing_counts.append(len(page.get_drawings()))
                except Exception:
                    drawing_counts.append(0)
        common = {x for x, n in xref_pages.items() if count >= 4 and n >= 0.6 * count}
        baseline = statistics.median(drawing_counts) if drawing_counts else 0
        ocr_left = _ocr_budget(apple=200, tesseract=60)
        parts: List[str] = []
        visual: List[str] = []
        empty_pages = 0
        ocr_pages = 0
        for i, page in enumerate(doc):
            number = i + 1
            text = _clean_text(_page_text(page))
            reasons: Set[str] = set()
            area = max(1.0, page.rect.width * page.rect.height)
            image_area = 0.0
            try:
                for info in page.get_image_info(xrefs=True):
                    if info.get("xref") in common:
                        continue
                    x0, y0, x1, y1 = info.get("bbox", (0, 0, 0, 0))
                    image_area += max(0.0, x1 - x0) * max(0.0, y1 - y0)
            except Exception:
                pass
            if image_area > 0.10 * area and not (image_area > 0.85 * area and len(text) > 300):
                reasons.add("figure or picture")
            if not long_doc and drawing_counts and drawing_counts[i] > baseline + 25:
                reasons.add("chart or diagram")
            if _mathy(text):
                reasons.add("equations")
            needs_text = len(text) < 25 and (image_area > 0 or not text)
            has_figure = bool(reasons & {"figure or picture", "chart or diagram"})
            figure_labels: List[str] = []
            if ocr_left > ocr_pages and (needs_text or has_figure):
                # Read the page as a picture: the whole text of a scan, or the labels inside a figure.
                try:
                    png = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False).tobytes("png")
                    lines_found = _ocr.read_lines(data=png, mode="page" if needs_text else "labels")
                    ocr_pages += 1
                    if needs_text:
                        text = _clean_text("\n".join(ln.text for ln in lines_found)) or text
                        if len(text) >= 25:
                            reasons.add("scanned page (text read from the image)")
                    else:
                        figure_labels = _ocr.labels_not_in(lines_found, text)
                except Exception:
                    pass
            if len(text) < 25:
                empty_pages += 1
                reasons.add("no text layer (scan or image-only page)")
            tables: List[str] = []
            if not long_doc and len(text) > 40:
                try:
                    for table in page.find_tables().tables[:4]:
                        md = table.to_markdown(clean=True) if hasattr(table, "to_markdown") else ""
                        if md and md.count("\n") >= 2:
                            tables.append(md.strip())
                except Exception:
                    pass
            block = [f"## [Page {number}]", ""]
            if reasons:
                visual.append(f"Page {number}")
                block.append(f"> Visual content: {', '.join(sorted(reasons))}. View the page image for exact details.\n")
            block.append(text or "(no extractable text on this page)")
            if figure_labels:
                block += ["", "Text in the figure (read from the image): " + "; ".join(_safe_inline(x) for x in figure_labels)]
            for table in tables:
                block += ["", "Detected table:", "", table]
            parts.append("\n".join(block))
    notes = []
    if empty_pages:
        notes.append(f"{empty_pages} of {count} pages have no text layer (scanned or image-only)"
                     + ("; OCR was applied where possible" if ocr_pages else "; view those pages as images"))
    return Extracted("pdf", "\n\n".join(parts), units=f"{count} pages", notes=notes, visual_units=visual)


# ---------------------------------------------------------------- PowerPoint

def _alt_text(shape) -> str:
    try:
        for el in shape._element.iter():
            if el.tag.endswith("}cNvPr"):
                return (el.get("descr") or el.get("title") or "").strip()
    except Exception:
        pass
    return ""


def _shape_type(shape):
    try:
        return shape.shape_type
    except Exception:
        return None


def _is_picture(shape) -> bool:
    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        if _shape_type(shape) == MSO_SHAPE_TYPE.PICTURE:
            return True
    except Exception:
        pass
    return shape.__class__.__name__ in ("Picture", "PlaceholderPicture")


MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
A14 = "{http://schemas.microsoft.com/office/drawing/2010/main}"
DML = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


class _AltItem:
    """A shape PowerPoint wraps in mc:AlternateContent (equation text boxes, ink, 3-D models, some SVGs).
    python-pptx skips these entirely, so read the newer form (Choice), with the fallback picture as backup."""

    def __init__(self, el):
        from .omml import math_text

        choice = el.find(f"{MC}Choice")
        fallback = el.find(f"{MC}Fallback")
        self.lines: List[str] = []
        self.math = False
        self.blip = None
        self.top = self.left = 0
        self.cx = self.cy = 0
        for branch in (choice, fallback):
            if branch is None:
                continue
            off = branch.find(f".//{DML}off")
            ext = branch.find(f".//{DML}ext")
            if off is not None and not (self.top or self.left):
                self.left, self.top = int(off.get("x", 0)), int(off.get("y", 0))
            if ext is not None and not self.cx:
                self.cx, self.cy = int(ext.get("cx", 0)), int(ext.get("cy", 0))
            if not self.lines:
                for para in branch.iter(f"{DML}p"):
                    pieces = []
                    for node in para:
                        if node.tag in (f"{DML}r", f"{DML}fld"):
                            pieces.append("".join(t.text or "" for t in node.iter(f"{DML}t")))
                        elif node.tag == f"{A14}m":
                            eq = math_text(node)
                            if eq:
                                pieces.append(f" {eq} ")
                                self.math = True
                    text = re.sub(r"\s+", " ", "".join(pieces)).strip()
                    if text:
                        self.lines.append("- " + text)
            if self.blip is None:
                blip = branch.find(f".//{DML}blip")
                if blip is not None and blip.get(f"{REL}embed"):
                    self.blip = blip.get(f"{REL}embed")


def _alt_items(shapes) -> list:
    container = getattr(shapes, "_spTree", None)
    if container is None:
        container = getattr(shapes, "_grpSp", None)
    if container is None:
        return []
    items = []
    for el in container.iterchildren(f"{MC}AlternateContent"):
        try:
            items.append(_AltItem(el))
        except Exception:
            continue
    return items


def _ordered(shapes) -> list:
    def key(s):
        return (getattr(s, "top", None) or 0, getattr(s, "left", None) or 0)

    try:
        return sorted(list(shapes), key=key)
    except Exception:
        return list(shapes)


def _text_lines(frame) -> List[str]:
    lines: List[str] = []
    for para in frame.paragraphs:
        text = "".join(run.text for run in para.runs) or para.text or ""
        text = re.sub(r"\s+", " ", text.replace("\x0b", " ")).strip()
        if not text:
            continue
        level = getattr(para, "level", 0) or 0
        lines.append(("  " * level) + "- " + text)
    return lines


def _chart_lines(chart) -> List[str]:
    out: List[str] = []
    try:
        ctype = str(chart.chart_type).split(" (")[0].split(".")[-1].replace("_", " ").lower()
    except Exception:
        ctype = "chart"
    title = ""
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            title = chart.chart_title.text_frame.text.strip()
    except Exception:
        pass
    out.append(f"Chart ({ctype}){': ' + title if title else ''}")
    num = lambda v: f"{v:.6g}" if isinstance(v, float) else ("" if v is None else str(v))  # noqa: E731
    try:
        if "xy" in ctype or "scatter" in ctype or "bubble" in ctype:
            rows = []
            for p in chart.plots:
                for ser in p.series:
                    xs, ys = _xy_points(ser)
                    for x, y in list(zip(xs, ys))[:200]:
                        rows.append([str(ser.name), num(x), num(y)])
            if rows:
                out += ["", _md_table(rows, ["Series", "x", "y"])]
            return out
        plot = chart.plots[0]
        cats = [str(c) for c in plot.categories]
        series = [(s.name, list(s.values)) for p in chart.plots for s in p.series]
        if cats and series:
            header = ["Category"] + [str(name) for name, _ in series]
            rows = []
            for idx, cat in enumerate(cats[:200]):
                row = [cat]
                for _, values in series:
                    row.append(num(values[idx]) if idx < len(values) else "")
                rows.append(row)
            out += ["", _md_table(rows, header)]
            if len(cats) > 200:
                out.append(f"(First 200 of {len(cats)} categories.)")
    except Exception:
        pass
    return out


def _xy_points(series):
    """x and y values of a scatter/bubble series, read from the chart's own data cache."""
    C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"

    def values(tag: str) -> list:
        node = series._element.find(f"{C}{tag}")
        if node is None:
            return []
        pts = {}
        for pt in node.iter(f"{C}pt"):
            v = pt.find(f"{C}v")
            if v is not None and v.text is not None:
                try:
                    pts[int(pt.get("idx", len(pts)))] = float(v.text)
                except ValueError:
                    pts[int(pt.get("idx", len(pts)))] = v.text
        return [pts[k] for k in sorted(pts)]
    return values("xVal"), values("yVal")


def _smartart_text(slide) -> List[str]:
    texts: List[str] = []
    try:
        for rel in slide.part.rels.values():
            if rel.reltype.endswith("/diagramData"):
                blob = rel.target_part.blob
                found = re.findall(rb"<a:t>([^<]*)</a:t>", blob)
                texts += [re.sub(r"\s+", " ", t.decode("utf-8", "replace")).strip() for t in found if t.strip()]
    except Exception:
        pass
    import html as _html

    return [_html.unescape(t) for t in texts if t]


def _region(point, box) -> str:
    x0, y0, x1, y1 = box
    rx = (point[0] - x0) / max(1, x1 - x0)
    ry = (point[1] - y0) / max(1, y1 - y0)
    if not (-0.02 <= rx <= 1.02 and -0.02 <= ry <= 1.02):
        side = ("left of it" if rx < 0 else "right of it" if rx > 1 else "above it" if ry < 0 else "below it")
        return side
    h = "left" if rx < 0.33 else "right" if rx > 0.67 else ""
    v = "upper" if ry < 0.33 else "lower" if ry > 0.67 else "middle"
    return (f"{v} {h}".strip() if h or v != "middle" else "center").replace("middle ", "")


def _add_picture_labels(slide, title_shape, pic_records, lines: List[str]) -> None:
    """Text boxes and arrows the instructor laid over a picture: say which label points where."""
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

    def bbox(sh):
        if None in (sh.left, sh.top, sh.width, sh.height):
            return None
        return (sh.left, sh.top, sh.left + sh.width, sh.top + sh.height)

    labels, arrows, pics = [], [], {}
    for sh in slide.shapes:
        try:
            box = bbox(sh)
            if box is None or (title_shape is not None and sh.shape_id == title_shape.shape_id):
                continue
            geom = sh._element.find(f".//{A}prstGeom")
            prst = geom.get("prst") if geom is not None else ""
            if sh._element.tag.endswith("}cxnSp") or prst in ("line", "straightConnector1"):
                xfrm = sh._element.find(f".//{A}xfrm")
                bx, by, ex, ey = box
                if xfrm is not None and xfrm.get("flipH") == "1":
                    bx, ex = ex, bx
                if xfrm is not None and xfrm.get("flipV") == "1":
                    by, ey = ey, by
                arrows.append(((bx, by), (ex, ey)))
            elif _is_picture(sh):
                pics[sh.shape_id] = box
            elif getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                text = re.sub(r"\s+", " ", sh.text_frame.text).strip()
                if text and len(text) <= 60 and len(text.split()) <= 6:
                    labels.append((text, box))
        except Exception:
            continue
    tol = 137160   # 0.15 inch
    for shape_id, at in sorted(pic_records, key=lambda r: -r[1]):
        pb = pics.get(shape_id)
        if pb is None:
            continue
        w, h = pb[2] - pb[0], pb[3] - pb[1]
        near_box = (pb[0] - 0.25 * w, pb[1] - 0.25 * h, pb[2] + 0.25 * w, pb[3] + 0.25 * h)
        found = []
        for text, lb in labels:
            pointer = None
            for a, b in arrows:
                for start, tip in ((a, b), (b, a)):
                    near_label = lb[0] - tol <= start[0] <= lb[2] + tol and lb[1] - tol <= start[1] <= lb[3] + tol
                    in_pic = pb[0] <= tip[0] <= pb[2] and pb[1] <= tip[1] <= pb[3]
                    if near_label and in_pic:
                        pointer = tip
            center = ((lb[0] + lb[2]) / 2, (lb[1] + lb[3]) / 2)
            if pointer is not None:
                found.append(f"{_safe_inline(text)} (arrow to the {_region(pointer, pb)} of the picture)")
            elif near_box[0] <= center[0] <= near_box[2] and near_box[1] <= center[1] <= near_box[3]:
                where = _region(center, pb)
                found.append(f"{_safe_inline(text)} ({where if 'of it' in where else 'on the picture, ' + where})")
        if found:
            lines.insert(at, "Labels on this picture: " + "; ".join(found))


def _pptx(path: Path) -> Extracted:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    try:
        prs = Presentation(str(path))
    except Exception as exc:
        return Extracted("slides", "", notes=[f"Couldn't open this deck ({exc.__class__.__name__}). Open the original."])
    slide_area = max(1, (prs.slide_width or 1) * (prs.slide_height or 1))
    slides = list(prs.slides)                      # python-pptx follows presentation order (sldIdLst)
    assets = assets_dir_for(path)
    _clear_assets(assets)

    def pictures(shapes):
        for shape in shapes:
            if _shape_type(shape) == MSO_SHAPE_TYPE.GROUP:
                yield from pictures(shape.shapes)
            elif _is_picture(shape):
                yield shape

    freq: Counter = Counter()
    for slide in slides:
        seen = set()
        for pic in pictures(slide.shapes):
            try:
                seen.add(hashlib.sha1(pic.image.blob).hexdigest())
            except Exception:
                continue
        freq.update(seen)
    common = {h for h, n in freq.items() if len(slides) >= 4 and n >= max(3, 0.5 * len(slides))}

    parts: List[str] = []
    visual: List[str] = []
    saved: List[str] = []
    ocr_left = [_ocr_budget(apple=400, tesseract=80)]   # pictures to read per deck
    for idx, slide in enumerate(slides, 1):
        title_shape = None
        try:
            title_shape = slide.shapes.title
        except Exception:
            pass
        title = ""
        if title_shape is not None and title_shape.has_text_frame:
            title = re.sub(r"\s+", " ", title_shape.text_frame.text).strip()
        hidden = slide._element.get("show") == "0"
        lines: List[str] = []
        reasons: Set[str] = set()
        pic_count = 0
        pic_records: List[tuple] = []     # (shape id, where its lines end) for labels laid over pictures

        def save_picture(blob: bytes, ext: str, alt: str, shape_id=None) -> None:
            nonlocal pic_count
            if ext == ".jpeg":
                ext = ".jpg"
            if ext not in VIEWABLE_IMAGE_EXT:
                converted = _to_png(blob) if ext in (".tif", ".tiff", ".bmp") else None
                if converted is None:
                    lines.append(f"[image ({ext.lstrip('.')} format, open the original to view){': ' + alt if alt else ''}]")
                    reasons.add("figure or picture")
                    return
                blob, ext = converted, ".png"
            pic_count += 1
            rel_path = _save_asset(assets, f"slide-{idx:02d}-{pic_count}{ext}", blob)
            saved.append(rel_path)
            reasons.add("figure or picture")
            lines.append(md_link(f"Slide {idx} image {pic_count}{': ' + alt if alt else ''}", rel_path))
            if ocr_left[0] > 0:
                ocr_left[0] -= 1
                try:
                    found = _ocr.summary(_ocr.read_lines(data=blob, mode="labels"))
                except Exception:
                    found = ""
                if found:
                    lines.append("Text in the image: " + _safe_inline(found))
            if shape_id is not None:
                pic_records.append((shape_id, len(lines)))

        def small_picture(blob: bytes, ext: str, alt: str) -> None:
            """Pictures too small to be a figure: often a formula or a label. Keep them if they hold text."""
            if _tiny_image(blob):
                return
            found = ""
            if ocr_left[0] > 0:
                ocr_left[0] -= 1
                try:
                    found = _ocr.summary(_ocr.read_lines(data=blob, mode="labels"))
                except Exception:
                    found = ""
            elif not _ocr.available():
                found = "?"
            if not found and not alt:
                return
            before = len(reasons)
            save_picture(blob, ext, alt)
            if found and found != "?":
                reasons.add("small picture with text (a formula or label)")
            elif len(reasons) == before:
                reasons.add("small picture")

        def walk(shapes):
            for shape in _ordered(list(shapes) + _alt_items(shapes)):
                if isinstance(shape, _AltItem):
                    lines.extend(shape.lines)
                    if shape.math:
                        reasons.add("equations")
                    elif not shape.lines and shape.blip:
                        try:
                            part = slide.part.related_part(shape.blip)
                            blob = part.blob
                            ext = "." + (getattr(part, "ext", "") or "png").lower().lstrip(".")
                        except Exception:
                            continue
                        if shape.cx * shape.cy >= 0.03 * slide_area:
                            save_picture(blob, ext, "")
                        else:
                            small_picture(blob, ext, "")
                    continue
                if title_shape is not None and shape.shape_id == title_shape.shape_id:
                    continue
                stype = _shape_type(shape)
                if stype == MSO_SHAPE_TYPE.GROUP:
                    walk(shape.shapes)
                    continue
                if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                    lines.extend(_text_lines(shape.text_frame))
                if getattr(shape, "has_table", False) and shape.has_table:
                    rows = [[c.text for c in row.cells] for row in shape.table.rows]
                    if rows:
                        lines.extend(["", _md_table(rows), ""])
                if getattr(shape, "has_chart", False) and shape.has_chart:
                    reasons.add("chart")
                    lines.extend([""] + _chart_lines(shape.chart) + [""])
                if _is_picture(shape):
                    alt = _alt_text(shape)
                    try:
                        blob = shape.image.blob
                        ext = "." + (shape.image.ext or "png").lower()
                        digest = hashlib.sha1(blob).hexdigest()
                    except Exception:
                        if alt:
                            lines.append(f"[image: {alt}]")
                        continue
                    area = (shape.width or 0) * (shape.height or 0)
                    if digest in common:          # template decoration repeated on most slides
                        continue
                    if area < 0.03 * slide_area:
                        small_picture(blob, ext, alt)
                        continue
                    save_picture(blob, ext, alt, shape.shape_id)

        walk(slide.shapes)
        if pic_records:
            _add_picture_labels(slide, title_shape, pic_records, lines)
        smart = _smartart_text(slide)
        if smart:
            reasons.add("diagram (SmartArt)")
            lines.append("Diagram: " + " / ".join(smart))
        if _mathy("\n".join(lines)):
            reasons.add("equations")
        notes = ""
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            notes = ""
        header = f"## [Slide {idx}]" + (f" {title}" if title else "") + (" (hidden slide)" if hidden else "")
        block = [header, ""]
        if reasons:
            visual.append(f"Slide {idx}")
            block.append(f"> Visual content: {', '.join(sorted(reasons))}. View the slide image for exact details.\n")
        block.append("\n".join(lines).strip() or "(no text on this slide)")
        if notes:
            block += ["", "Speaker notes:", "", _clean_text(notes)]
        parts.append("\n".join(block))
    notes_out = []
    if saved:
        notes_out.append(f"{len(saved)} slide images saved in '{assets.name}/' (view them for charts and figures)")
    return Extracted("slides", "\n\n".join(parts), units=f"{len(slides)} slides", notes=notes_out,
                     visual_units=visual, assets=saved)


# ---------------------------------------------------------------- Word

def _docx(path: Path) -> Extracted:
    import mammoth

    assets = assets_dir_for(path)
    _clear_assets(assets)
    saved: List[str] = []
    counter = [0]
    seen: dict = {}
    types = {"image/png": ".png", "image/jpeg": ".jpg", "image/jpg": ".jpg", "image/gif": ".gif", "image/webp": ".webp"}

    def convert_image(image):
        alt = (getattr(image, "alt_text", None) or "").strip()
        try:
            with image.open() as fh:
                data = fh.read()
        except Exception:
            return {"src": "", "alt": alt or "image"}
        ext = types.get((image.content_type or "").lower())
        if not ext:
            converted = _to_png(data)
            if converted is None:
                return {"src": "", "alt": alt or "image"}
            data, ext = converted, ".png"
        digest = hashlib.sha1(data).hexdigest()
        if digest in seen:                      # the same picture again (a logo, a repeated diagram)
            return {"src": seen[digest][0].replace(" ", "%20"), "alt": alt or seen[digest][1]}
        if _tiny_image(data):                  # bullets and icons; small equation images are kept
            return {"src": "", "alt": alt or "icon"}
        counter[0] += 1
        rel_path = _save_asset(assets, f"image-{counter[0]:02d}{ext}", data)
        saved.append(rel_path)
        label = alt or f"figure {counter[0]}"
        seen[digest] = (rel_path, label)
        if counter[0] <= _ocr_budget(apple=150, tesseract=40):
            try:
                found = _safe_inline(_ocr.summary(_ocr.read_lines(data=data, mode="labels"), limit=40))[:500]
            except Exception:
                found = ""
            if found:
                label += f" (text in the image: {found})"
        return {"src": rel_path.replace(" ", "%20"), "alt": label}

    from .omml import cleanup, docx_with_linear_math

    with_math = docx_with_linear_math(path)    # equations kept in place as [equation: ...] text
    try:
        with open(with_math or path, "rb") as fh:
            result = mammoth.convert_to_html(fh, convert_image=mammoth.images.img_element(convert_image))
    except Exception as exc:
        return Extracted("document", "", notes=[f"Couldn't read this Word file ({exc.__class__.__name__})."])
    finally:
        cleanup(with_math)
    md = simple_html_to_markdown(result.value)
    notes = [f"{len(saved)} figures saved in '{assets.name}/'"] if saved else []
    if with_math is not None:
        notes.append("Equations are written in linear form as [equation: ...]; open the original for the typeset layout")
    elif "<m:oMath" in _docx_xml(path):
        notes.append("Contains Word equations that couldn't be converted; open the original for formulas")
    return Extracted("document", md, units=f"{len(md.split())} words", notes=notes,
                     visual_units=["Figures"] if saved else [], assets=saved)


def _tiny_image(data: bytes) -> bool:
    """Bullets, icons and spacer images: too small to carry content."""
    try:
        from PIL import Image

        with Image.open(io.BytesIO(data)) as im:
            w, h = im.size
        return w < 24 or h < 12 or (w * h) < 900
    except Exception:
        return len(data) < 300


def _docx_xml(path: Path) -> str:
    try:
        import zipfile

        with zipfile.ZipFile(path) as z:
            return z.read("word/document.xml").decode("utf-8", "replace")
    except Exception:
        return ""


# ---------------------------------------------------------------- spreadsheets

def _fmt_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.10g}"
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat(sep=" ") if hasattr(value, "hour") else value.isoformat()
        except TypeError:
            return value.isoformat()
    return str(value)


def _xlsx(path: Path) -> Extracted:
    import openpyxl
    from openpyxl.utils import get_column_letter

    try:
        values = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:
        return Extracted("spreadsheet", "", notes=[f"Couldn't open this workbook ({exc.__class__.__name__})."])
    try:
        formulas = openpyxl.load_workbook(str(path), read_only=True, data_only=False)
    except Exception:
        formulas = None
    parts: List[str] = []
    notes: List[str] = []
    for ws in values.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True, max_row=MAX_SHEET_ROWS)]
        while rows and not any(v not in (None, "") for v in rows[-1]):
            rows.pop()
        width = 0
        for r in rows:
            for j in range(len(r) - 1, -1, -1):
                if r[j] not in (None, ""):
                    width = max(width, j + 1)
                    break
        full_width = width
        width = min(width, MAX_SHEET_COLS)
        block = [f"## [Sheet: {ws.title}]", ""]
        if not rows or not width:
            block.append("(empty sheet)")
            parts.append("\n".join(block))
            continue
        header = ["Row"] + [get_column_letter(j + 1) for j in range(width)]
        table_rows = [[str(i + 1)] + [_fmt_cell(v) for v in r[:width]] for i, r in enumerate(rows)
                      if any(v not in (None, "") for v in r[:width])]
        block.append(_md_table(table_rows, header))
        if full_width > width:
            block.append(f"\n(Showing the first {width} of {full_width} columns; the original file has the rest.)")
            notes.append(f"Sheet '{ws.title}' shows {width} of {full_width} columns")
        max_row = getattr(ws, "max_row", None)
        if max_row and max_row > MAX_SHEET_ROWS:
            block.append(f"\n(Showing the first {MAX_SHEET_ROWS} of {max_row} rows.)")
            notes.append(f"Sheet '{ws.title}' truncated to {MAX_SHEET_ROWS} rows")
        if formulas is not None:
            found = []
            try:
                for row in formulas[ws.title].iter_rows(max_row=MAX_SHEET_ROWS, max_col=MAX_SHEET_COLS):
                    for c in row:
                        v = getattr(c, "value", None)
                        if isinstance(v, str) and v.startswith("="):
                            found.append(f"- {c.coordinate}: `{v}`")
                            if len(found) >= 300:
                                break
                    if len(found) >= 300:
                        break
            except Exception:
                pass
            if found:
                block += ["", "Formulas (cell: formula):", ""] + found
        parts.append("\n".join(block))
    values.close()
    if formulas is not None:
        formulas.close()
    return Extracted("spreadsheet", "\n\n".join(parts), units=f"{len(parts)} sheets", notes=notes)


def _csv(path: Path) -> Extracted:
    raw = path.read_bytes()[:MAX_TEXT_BYTES].decode("utf-8-sig", "replace")
    try:
        dialect = csv.Sniffer().sniff(raw[:4096], delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
    rows = list(csv.reader(io.StringIO(raw), dialect))
    total = len(rows)
    shown = rows[: MAX_SHEET_ROWS + 1]
    body = "## [Sheet: data]\n\n" + (_md_table([r[:MAX_SHEET_COLS] for r in shown]) if shown else "(empty)")
    cols = max((len(r) for r in shown), default=0)
    if cols > MAX_SHEET_COLS:
        body += f"\n\n(Showing the first {MAX_SHEET_COLS} of {cols} columns; the original file has the rest.)"
    if total > len(shown):
        body += f"\n\n(Showing {len(shown) - 1} of {total - 1} data rows.)"
    return Extracted("spreadsheet", body, units=f"{max(0, total - 1)} rows")


# ---------------------------------------------------------------- other text formats

def _notebook(path: Path) -> Extracted:
    try:
        nb = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return _text(path)
    lang = ((nb.get("metadata") or {}).get("kernelspec") or {}).get("language") or "python"
    parts: List[str] = []
    for i, cell in enumerate(nb.get("cells") or [], 1):
        src = "".join(cell.get("source") or [])
        if not src.strip():
            continue
        if cell.get("cell_type") == "markdown":
            parts.append(f"## [Cell {i}]\n\n{src.strip()}")
            continue
        block = f"## [Cell {i}] code\n\n```{lang}\n{src.rstrip()}\n```"
        outputs = []
        for out in cell.get("outputs") or []:
            text = "".join(out.get("text") or (out.get("data") or {}).get("text/plain") or [])
            if text.strip():
                outputs.append(text.strip()[:3000])
        if outputs:
            block += "\n\nOutput:\n\n```\n" + "\n".join(outputs) + "\n```"
        parts.append(block)
    return Extracted("notebook", "\n\n".join(parts), units=f"{len(parts)} cells")


def _html_file(path: Path) -> Extracted:
    raw = path.read_bytes()[:MAX_TEXT_BYTES].decode("utf-8", "replace")
    md = simple_html_to_markdown(raw)
    return Extracted("web", md, units=f"{len(md.split())} words")


def _text(path: Path) -> Extracted:
    raw = path.read_bytes()
    truncated = len(raw) > MAX_TEXT_BYTES
    text = raw[:MAX_TEXT_BYTES].decode("utf-8", "replace")
    lang = CODE_LANG.get(path.suffix.lower())
    if lang:
        body = f"```{lang}\n{text.rstrip()}\n```"
    else:
        body = _clean_text(text)
    notes = ["Large file: only the first 2 MB was indexed"] if truncated else []
    return Extracted("text", body, units=f"{len(text.splitlines())} lines", notes=notes)


def _image(path: Path) -> Extracted:
    size = ""
    try:
        from PIL import Image

        with Image.open(path) as img:
            size = f"{img.width}x{img.height} px"
    except Exception:
        pass
    text = ocr_image_file(path) if path.suffix.lower() not in (".svg",) else ""
    body = "## [Image]\n\n> Visual content: this file is an image. View it directly.\n"
    if text:
        body += "\nText found in the image (OCR):\n\n" + text
    viewable = path.suffix.lower() in VIEWABLE_IMAGE_EXT
    notes = [] if viewable else ["Format may need converting before an AI can view it"]
    return Extracted("image", body, units=size, notes=notes, visual_units=["Image"])


def _legacy(path: Path) -> Extracted:
    ext = path.suffix.lower()
    work = Path(tempfile.mkdtemp(prefix="ss-conv-"))
    try:
        target = {".ppt": ".pptx", ".pps": ".pptx", ".ppsx": ".pptx", ".odp": ".pptx", ".key": ".pptx",
                  ".xls": ".xlsx", ".ods": ".xlsx", ".numbers": ".xlsx"}.get(ext, ".docx")
        converted = convert_with_soffice(path, target, work)
        if converted:
            inner = extract(converted)
            inner.notes.append(f"Converted from {ext} with LibreOffice")
            if inner.assets:  # move saved images next to the original
                src_assets = assets_dir_for(converted)
                dst_assets = assets_dir_for(path)
                _clear_assets(dst_assets)
                if src_assets.exists():
                    shutil.move(str(src_assets), str(dst_assets))
                inner.assets = [a.replace(src_assets.name, dst_assets.name, 1) for a in inner.assets]
                inner.body = inner.body.replace(md_link("x", src_assets.name)[4:-1], md_link("x", dst_assets.name)[4:-1])
            return inner
        if ext == ".xls":
            try:
                import xlrd

                book = xlrd.open_workbook(str(path))
                parts = []
                for sheet in book.sheets():
                    rows = [[_fmt_cell(sheet.cell_value(r, c)) for c in range(min(sheet.ncols, MAX_SHEET_COLS))]
                            for r in range(min(sheet.nrows, MAX_SHEET_ROWS))]
                    parts.append(f"## [Sheet: {sheet.name}]\n\n" + (_md_table(rows) if rows else "(empty sheet)"))
                return Extracted("spreadsheet", "\n\n".join(parts), units=f"{len(parts)} sheets")
            except Exception:
                pass
        if shutil.which("textutil") and ext in (".doc", ".rtf", ".odt", ".wpd"):
            out = subprocess.run(["textutil", "-convert", "txt", "-stdout", str(path)],
                                 capture_output=True, text=True, timeout=120)
            if out.returncode == 0 and out.stdout.strip():
                return Extracted("document", _clean_text(out.stdout), units=f"{len(out.stdout.split())} words",
                                 notes=["Converted with macOS textutil"])
        return Extracted("unsupported", "", notes=[
            f"'{ext}' files need LibreOffice to be read (free: brew install --cask libreoffice). Open the original meanwhile."])
    finally:
        shutil.rmtree(work, ignore_errors=True)
