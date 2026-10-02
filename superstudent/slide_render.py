"""Draw a PowerPoint slide as a picture, without LibreOffice.

The AI needs to see a slide as a whole: a diagram with the labels, arrows and text boxes the instructor laid
over it, not just the embedded photo on its own. LibreOffice renders slides exactly; when it isn't installed,
this draws them from the file itself: pictures, text, shapes, lines and arrows, tables and groups, each in its
place. Fonts, theme colors and effects are approximate; positions and content are faithful.
"""

from __future__ import annotations

import io
import math
import re
from functools import lru_cache
from pathlib import Path
from typing import Callable, List, Optional, Tuple

from PIL import Image, ImageDraw, ImageFont

EMU_PER_PT = 12700
FONTS = {
    False: ["/System/Library/Fonts/Supplemental/Arial.ttf", "/Library/Fonts/Arial.ttf",
            "/System/Library/Fonts/Helvetica.ttc", "C:/Windows/Fonts/arial.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/dejavu/DejaVuSans.ttf"],
    True: ["/System/Library/Fonts/Supplemental/Arial Bold.ttf", "/Library/Fonts/Arial Bold.ttf",
           "/System/Library/Fonts/Helvetica.ttc", "C:/Windows/Fonts/arialbd.ttf",
           "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf"],
}
OFFICE_BLUE = (68, 114, 196)
INK = (34, 34, 34)
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

Transform = Callable[[float, float], Tuple[float, float]]


@lru_cache(maxsize=64)
def _font(size: int, bold: bool = False):
    size = max(6, int(size))
    for path in FONTS[bold] + FONTS[False]:
        try:
            return ImageFont.truetype(path, size)
        except OSError:
            continue
    try:
        return ImageFont.load_default(size=size)
    except TypeError:
        return ImageFont.load_default()


def _luma(rgb: Tuple[int, int, int]) -> float:
    r, g, b = rgb
    return (0.299 * r + 0.587 * g + 0.114 * b) / 255


def _rgb(color_format) -> Optional[Tuple[int, int, int]]:
    """An explicit RGB color, or None when it's a theme color we can't resolve."""
    try:
        value = color_format.rgb
        if value is None:
            return None
        text = str(value)
        return (int(text[0:2], 16), int(text[2:4], 16), int(text[4:6], 16))
    except Exception:
        return None


def _solid_fill(fill) -> Optional[Tuple[int, int, int]]:
    try:
        from pptx.enum.dml import MSO_FILL

        if fill.type == MSO_FILL.SOLID:
            return _rgb(fill.fore_color) or OFFICE_BLUE
    except Exception:
        pass
    return None


def _no_fill(fill) -> bool:
    try:
        from pptx.enum.dml import MSO_FILL

        return fill.type == MSO_FILL.BACKGROUND
    except Exception:
        return False


class SlidePainter:
    def __init__(self, prs, slide, long_edge: int = 1600):
        self.W = prs.slide_width or 9144000
        self.H = prs.slide_height or 6858000
        self.k = long_edge / max(self.W, self.H)
        self.bg = self._background(slide)
        self.img = Image.new("RGB", (max(1, round(self.W * self.k)), max(1, round(self.H * self.k))), self.bg)
        self.draw = ImageDraw.Draw(self.img)
        self.slide = slide

    # ---------------------------------------------------------------- helpers
    def _background(self, slide) -> Tuple[int, int, int]:
        for owner in (slide, getattr(slide, "slide_layout", None),
                      getattr(getattr(slide, "slide_layout", None), "slide_master", None)):
            if owner is None:
                continue
            try:
                color = _solid_fill(owner.background.fill)
                if color:
                    return color
            except Exception:
                continue
        return (255, 255, 255)

    def px(self, v: float) -> float:
        return v * self.k

    def box(self, shape, tf: Transform) -> Optional[Tuple[float, float, float, float]]:
        try:
            left, top, width, height = shape.left, shape.top, shape.width, shape.height
        except Exception:
            return None
        if left is None or top is None or width is None or height is None:
            return None
        x0, y0 = tf(left, top)
        x1, y1 = tf(left + width, top + height)
        return (self.px(min(x0, x1)), self.px(min(y0, y1)), self.px(max(x0, x1)), self.px(max(y0, y1)))

    # ---------------------------------------------------------------- shapes
    def paint(self) -> Image.Image:
        self.walk(self.slide.shapes, lambda x, y: (x, y))
        return self.img

    def walk(self, shapes, tf: Transform) -> None:
        container = getattr(shapes, "_spTree", None)
        if container is None:
            container = getattr(shapes, "_grpSp", None)
        shape_list = list(shapes)
        if container is None:
            children = shape_list
        else:
            by_element = {sh._element: sh for sh in shape_list}
            # document order, including mc:AlternateContent (equation boxes, ink) that python-pptx skips
            children = [by_element.get(ch, ch) for ch in container
                        if ch in by_element or ch.tag == f"{MC}AlternateContent"]
        for shape in children:
            try:
                if getattr(shape, "tag", None) == f"{MC}AlternateContent":
                    self.alternate(shape, tf)
                else:
                    self.shape(shape, tf)
            except Exception:
                continue   # one odd shape shouldn't lose the slide

    def alternate(self, el, tf: Transform) -> None:
        from .extract import _AltItem

        item = _AltItem(el)
        if not (item.cx and item.cy):
            return
        (x0, y0), (x1, y1) = tf(item.left, item.top), tf(item.left + item.cx, item.top + item.cy)
        box = (self.px(min(x0, x1)), self.px(min(y0, y1)), self.px(max(x0, x1)), self.px(max(y0, y1)))
        if item.lines:
            text = "\n".join(line[2:] if line.startswith("- ") else line for line in item.lines)
            color = (250, 250, 250) if _luma(self.bg) < 0.45 else INK
            font = _font(max(10, round(self.px(18 * EMU_PER_PT))))
            self.wrapped(text, (box[0] + 4, box[1] + 2, box[2] - 4, self.img.height), font, color)
        elif item.blip:
            try:
                blob = self.slide.part.related_part(item.blip).blob
                im = Image.open(io.BytesIO(blob))
                im.load()
            except Exception:
                return
            w, h = max(1, round(box[2] - box[0])), max(1, round(box[3] - box[1]))
            im = im.convert("RGBA").resize((w, h))
            self.img.paste(im, (round(box[0]), round(box[1])), im)

    def shape(self, shape, tf: Transform) -> None:
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        kind = None
        try:
            kind = shape.shape_type
        except Exception:
            pass
        if kind == MSO_SHAPE_TYPE.GROUP:
            self.walk(shape.shapes, self._group_tf(shape, tf))
            return
        tag = shape._element.tag
        if tag == f"{P}cxnSp" or kind == MSO_SHAPE_TYPE.LINE or self._prst(shape) in ("line", "straightConnector1"):
            self.line(shape, tf)
            return
        box = self.box(shape, tf)
        if box is None:
            return
        if shape.__class__.__name__ in ("Picture", "PlaceholderPicture") or kind == MSO_SHAPE_TYPE.PICTURE:
            self.picture(shape, box)
            return
        if getattr(shape, "has_table", False) and shape.has_table:
            self.table(shape, box)
            return
        if getattr(shape, "has_chart", False) and shape.has_chart:
            self.chart(shape, box)
            return
        if tag == f"{P}graphicFrame":
            self.placeholder_box(box, "Diagram (SmartArt): see the text version for its words")
            return
        self.autoshape(shape, box)

    def _prst(self, shape) -> str:
        geom = shape._element.find(f".//{A}prstGeom")
        return geom.get("prst") if geom is not None else ""

    def _group_tf(self, group, tf: Transform) -> Transform:
        xfrm = group._element.find(f"{P}grpSpPr/{A}xfrm")
        if xfrm is None:
            return tf
        def pair(tagname: str, a: str, b: str) -> Tuple[float, float]:
            el = xfrm.find(f"{A}{tagname}")
            return (float(el.get(a, 0)), float(el.get(b, 0))) if el is not None else (0.0, 0.0)
        ox, oy = pair("off", "x", "y")
        cx, cy = pair("ext", "cx", "cy")
        chx, chy = pair("chOff", "x", "y")
        chcx, chcy = pair("chExt", "cx", "cy")
        sx = cx / chcx if chcx else 1.0
        sy = cy / chcy if chcy else 1.0
        return lambda x, y: tf(ox + (x - chx) * sx, oy + (y - chy) * sy)

    def picture(self, shape, box) -> None:
        x0, y0, x1, y1 = box
        w, h = max(1, round(x1 - x0)), max(1, round(y1 - y0))
        try:
            im = Image.open(io.BytesIO(shape.image.blob))
            im.load()
        except Exception:
            self.placeholder_box(box, "Picture (format not shown here)")
            return
        try:
            cl, cr = float(shape.crop_left or 0), float(shape.crop_right or 0)
            ct, cb = float(shape.crop_top or 0), float(shape.crop_bottom or 0)
            if any((cl, cr, ct, cb)):
                W, H = im.size
                im = im.crop((max(0, round(W * cl)), max(0, round(H * ct)),
                              min(W, round(W * (1 - cr))), min(H, round(H * (1 - cb)))))
        except Exception:
            pass
        im = im.convert("RGBA").resize((w, h))
        rotation = float(getattr(shape, "rotation", 0) or 0)
        if rotation:
            im = im.rotate(-rotation, expand=True, resample=Image.BICUBIC)
        cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
        self.img.paste(im, (round(cx - im.width / 2), round(cy - im.height / 2)), im)

    def autoshape(self, shape, box) -> None:
        prst = self._prst(shape)
        fill_color = None
        outline = None
        try:
            fill_color = _solid_fill(shape.fill)
            if fill_color is None and not _no_fill(shape.fill) and self._styled(shape, "fillRef") and prst:
                fill_color = OFFICE_BLUE        # default shape style: theme accent fill
        except Exception:
            pass
        try:
            ln = shape.line
            if ln.fill.type is not None:
                outline = _rgb(ln.color) or (90, 90, 90)
            elif self._styled(shape, "lnRef") and prst:
                outline = (47, 82, 143)
        except Exception:
            pass
        x0, y0, x1, y1 = box
        width = max(1, round(self.px(self._line_width(shape))))
        if fill_color or outline:
            if prst in ("ellipse", "circle"):
                self.draw.ellipse(box, fill=fill_color, outline=outline, width=width)
            elif prst in ("roundRect", "snipRoundRect", "round2SameRect"):
                self.draw.rounded_rectangle(box, radius=min(x1 - x0, y1 - y0) * 0.15, fill=fill_color,
                                            outline=outline, width=width)
            elif "Arrow" in prst and prst not in ("curvedRightArrow",):
                self.arrow_shape(prst, box, fill_color or outline)
            elif prst in ("triangle", "rtTriangle"):
                self.draw.polygon([((x0 + x1) / 2, y0), (x1, y1), (x0, y1)], fill=fill_color, outline=outline)
            else:
                self.draw.rectangle(box, fill=fill_color, outline=outline, width=width)
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            is_box = shape._element.find(f"{P}nvSpPr/{P}cNvSpPr[@txBox='1']") is not None
            self.text(shape, box, fill_color or self.bg, anchor_default="t" if (is_box or not prst or prst == "rect" and not fill_color) else "ctr")

    def _styled(self, shape, ref: str) -> bool:
        el = shape._element.find(f"{P}style/{A}{ref}")
        return el is not None and el.get("idx", "0") != "0"

    def _line_width(self, shape) -> float:
        try:
            if shape.line.width:
                return float(shape.line.width)
        except Exception:
            pass
        return 12700 * 1.5

    def arrow_shape(self, prst: str, box, color) -> None:
        x0, y0, x1, y1 = box
        w, h = x1 - x0, y1 - y0
        if prst.startswith(("left",)):
            pts = [(x1, y0 + h * .3), (x0 + w * .4, y0 + h * .3), (x0 + w * .4, y0), (x0, y0 + h / 2),
                   (x0 + w * .4, y1), (x0 + w * .4, y0 + h * .7), (x1, y0 + h * .7)]
        elif prst.startswith(("up",)):
            pts = [(x0 + w * .3, y1), (x0 + w * .3, y0 + h * .4), (x0, y0 + h * .4), (x0 + w / 2, y0),
                   (x1, y0 + h * .4), (x0 + w * .7, y0 + h * .4), (x0 + w * .7, y1)]
        elif prst.startswith(("down",)):
            pts = [(x0 + w * .3, y0), (x0 + w * .3, y0 + h * .6), (x0, y0 + h * .6), (x0 + w / 2, y1),
                   (x1, y0 + h * .6), (x0 + w * .7, y0 + h * .6), (x0 + w * .7, y0)]
        else:
            pts = [(x0, y0 + h * .3), (x0 + w * .6, y0 + h * .3), (x0 + w * .6, y0), (x1, y0 + h / 2),
                   (x0 + w * .6, y1), (x0 + w * .6, y0 + h * .7), (x0, y0 + h * .7)]
        self.draw.polygon(pts, fill=color)

    def line(self, shape, tf: Transform) -> None:
        try:
            left, top, width, height = shape.left, shape.top, shape.width, shape.height
        except Exception:
            return
        xfrm = shape._element.find(f".//{A}xfrm")
        flip_h = xfrm is not None and xfrm.get("flipH") == "1"
        flip_v = xfrm is not None and xfrm.get("flipV") == "1"
        bx, by, ex, ey = left, top, left + width, top + height
        if flip_h:
            bx, ex = ex, bx
        if flip_v:
            by, ey = ey, by
        (bx, by), (ex, ey) = tf(bx, by), tf(ex, ey)
        begin, end = (self.px(bx), self.px(by)), (self.px(ex), self.px(ey))
        color = INK
        try:
            color = _rgb(shape.line.color) or color
        except Exception:
            pass
        width = max(2, round(self.px(self._line_width(shape))))
        self.draw.line([begin, end], fill=color, width=width)
        ln = shape._element.find(f".//{A}ln")
        if ln is not None:
            tail, head = ln.find(f"{A}tailEnd"), ln.find(f"{A}headEnd")
            if tail is not None and tail.get("type", "none") != "none":
                self.arrowhead(begin, end, color, width)
            if head is not None and head.get("type", "none") != "none":
                self.arrowhead(end, begin, color, width)

    def arrowhead(self, start, tip, color, width: int) -> None:
        angle = math.atan2(tip[1] - start[1], tip[0] - start[0])
        size = max(10, width * 4)
        left = (tip[0] - size * math.cos(angle - 0.45), tip[1] - size * math.sin(angle - 0.45))
        right = (tip[0] - size * math.cos(angle + 0.45), tip[1] - size * math.sin(angle + 0.45))
        self.draw.polygon([tip, left, right], fill=color)

    def table(self, shape, box) -> None:
        x0, y0, x1, y1 = box
        tbl = shape.table
        widths = [self.px(c.width) for c in tbl.columns]
        heights = [self.px(r.height) for r in tbl.rows]
        font = _font(max(9, min(22, round(min(heights or [20]) * 0.45))))
        y = y0
        for r, row in enumerate(tbl.rows):
            x = x0
            for c, cell in enumerate(row.cells):
                w = widths[c] if c < len(widths) else 60
                h = heights[r] if r < len(heights) else 20
                self.draw.rectangle([x, y, x + w, y + h], outline=(120, 120, 120),
                                    fill=(231, 236, 247) if r == 0 else None)
                self.wrapped(cell.text or "", (x + 4, y + 2, x + w - 4, y + h - 2), font, INK)
                x += w
            y += heights[r] if r < len(heights) else 20

    def chart(self, shape, box) -> None:
        title = ""
        try:
            chart = shape.chart
            if chart.has_title and chart.chart_title.has_text_frame:
                title = chart.chart_title.text_frame.text.strip()
            kind = str(chart.chart_type).split(" (")[0].split(".")[-1].replace("_", " ").lower()
        except Exception:
            kind = "chart"
        self.placeholder_box(box, f"Chart ({kind}){': ' + title if title else ''}. Its numbers are in the text version.")

    def placeholder_box(self, box, label: str) -> None:
        self.draw.rectangle(box, outline=(150, 150, 150), fill=(244, 244, 244))
        x0, y0, x1, y1 = box
        self.wrapped(label, (x0 + 8, y0 + 8, x1 - 8, y1 - 8), _font(max(10, round(self.px(14 * EMU_PER_PT)))), (90, 90, 90))

    # ---------------------------------------------------------------- text
    def text(self, shape, box, behind: Tuple[int, int, int], anchor_default: str = "t") -> None:
        frame = shape.text_frame
        paragraphs = []
        is_title = is_body = False
        try:
            if shape.is_placeholder:
                ptype = str(shape.placeholder_format.type)
                is_title = "TITLE" in ptype
                is_body = "BODY" in ptype or "OBJECT" in ptype
        except Exception:
            pass
        default_pt = 36 if is_title else (22 if is_body else 18)
        inherited_align, inherited_pt = self._inherited(shape, is_title, is_body)
        default_pt = inherited_pt or default_pt
        for para in frame.paragraphs:
            runs = list(para.runs)
            text = "".join(r.text for r in runs).replace("\x0b", "\n")
            if not text.strip():
                paragraphs.append(("", default_pt, False, None, 0, False, None))
                continue
            size = next((r.font.size for r in runs if r.font.size), None)
            pt = size.pt if size else default_pt
            bold = any(r.font.bold for r in runs)
            color = next((c for c in (_rgb(r.font.color) for r in runs if r.font.color and r.font.color.type) if c), None)
            level = getattr(para, "level", 0) or 0
            ppr = para._p.pPr
            bullet = False
            if ppr is not None and ppr.find(f"{A}buChar") is not None:
                bullet = True
            elif is_body and not (ppr is not None and ppr.find(f"{A}buNone") is not None):
                bullet = True
            align = inherited_align
            try:
                if para.alignment is not None:
                    align = str(para.alignment)
            except Exception:
                pass
            paragraphs.append((text, pt, bold, color, level, bullet, align))
        if not any(p[0].strip() for p in paragraphs):
            return
        body_pr = shape._element.find(f".//{A}bodyPr")
        anchor = (body_pr.get("anchor") if body_pr is not None else None) or anchor_default
        x0, y0, x1, y1 = box
        inset = self.px(91440)
        area = (x0 + inset, y0 + inset * 0.5, x1 - inset, y1 - inset * 0.5)
        dark_behind = _luma(behind) < 0.45
        scale = 1.0
        for _ in range(8):   # shrink text until it fits, like PowerPoint's autofit
            layout, height = self._layout(paragraphs, area, scale)
            if height <= (area[3] - area[1]) * 1.05 or scale < 0.45:
                break
            scale *= 0.88
        top = area[1]
        if anchor == "ctr":
            top = area[1] + max(0.0, ((area[3] - area[1]) - height) / 2)
        elif anchor == "b":
            top = area[3] - height
        y = top
        for line, font, color, indent, align, line_h in layout:
            if color is None or abs(_luma(color) - _luma(behind)) < 0.3:
                color = (250, 250, 250) if dark_behind else INK
            x = area[0] + indent
            if align and "CENTER" in align:
                x = area[0] + ((area[2] - area[0]) - font.getlength(line)) / 2
            elif align and "RIGHT" in align:
                x = area[2] - font.getlength(line)
            self.draw.text((x, y), line, fill=color, font=font)
            y += line_h

    def _inherited(self, shape, is_title: bool, is_body: bool) -> Tuple[Optional[str], Optional[float]]:
        """Alignment and size a placeholder inherits from its layout, then the slide master's text styles."""
        if not (is_title or is_body):
            return None, None
        align: Optional[str] = None
        size: Optional[float] = None
        names = {"ctr": "CENTER", "r": "RIGHT", "just": "JUSTIFY", "l": "LEFT"}
        sources = []
        try:
            layout_ph = self.slide.slide_layout.placeholders.get(idx=shape.placeholder_format.idx)
            if layout_ph is not None:
                sources.append(layout_ph._element.find(f".//{A}lstStyle/{A}lvl1pPr"))
                sources.append(layout_ph._element.find(f".//{A}p/{A}pPr"))
        except Exception:
            pass
        try:
            master = self.slide.slide_layout.slide_master._element
            style = "titleStyle" if is_title else "bodyStyle"
            sources.append(master.find(f".//{P}txStyles/{P}{style}/{A}lvl1pPr"))
        except Exception:
            pass
        for el in sources:
            if el is None:
                continue
            if align is None and el.get("algn"):
                align = names.get(el.get("algn"))
            rpr = el.find(f"{A}defRPr")
            if size is None and rpr is not None and rpr.get("sz"):
                size = int(rpr.get("sz")) / 100
        return align, size

    def _layout(self, paragraphs, area, scale: float):
        width = max(10.0, area[2] - area[0])
        lines = []
        total = 0.0
        for para in paragraphs:
            text, pt, bold, color, level, bullet, align = para
            size = max(7, round(self.px(pt * EMU_PER_PT) * scale))
            font = _font(size, bold)
            line_h = size * 1.2
            indent = self.px(level * 342900) * scale
            if not text.strip():
                lines.append(("", font, color, indent, align, line_h * 0.6))
                total += line_h * 0.6
                continue
            prefix = "• " if bullet else ""
            for chunk in text.split("\n"):
                for i, row in enumerate(self._wrap(prefix + chunk if chunk else chunk, font, width - indent)):
                    lines.append((row, font, color, indent + (font.getlength(prefix) if i and prefix else 0), align, line_h))
                    total += line_h
                prefix = ""
        return lines, total

    def _wrap(self, text: str, font, width: float) -> List[str]:
        words = text.split(" ")
        rows, cur = [], ""
        for word in words:
            trial = (cur + " " + word) if cur else word
            if font.getlength(trial) <= width or not cur:
                cur = trial
            else:
                rows.append(cur)
                cur = word
        rows.append(cur)
        return rows

    def wrapped(self, text: str, area, font, color) -> None:
        y = area[1]
        for para in re.split(r"\n+", text.strip()):
            for row in self._wrap(para, font, max(10, area[2] - area[0])):
                if y > area[3]:
                    return
                self.draw.text((area[0], y), row, fill=color, font=font)
                y += font.size * 1.2 if hasattr(font, "size") else 14


def render_slide(pptx_path: Path, number: int, out_path: Path, long_edge: int = 1600) -> Path:
    """Draw slide `number` (1-based) of a .pptx to a PNG."""
    from pptx import Presentation

    prs = Presentation(str(pptx_path))
    slides = list(prs.slides)
    if number < 1 or number > len(slides):
        raise ValueError(f"Slide {number} doesn't exist (this deck has {len(slides)}).")
    image = SlidePainter(prs, slides[number - 1], long_edge).paint()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    image.save(out_path, "PNG", optimize=True)
    return out_path


def slide_count(pptx_path: Path) -> int:
    from pptx import Presentation

    return len(Presentation(str(pptx_path)).slides)


def deck_to_pdf(pptx_path: Path, out_pdf: Path, long_edge: int = 1600) -> Path:
    """Every slide drawn into one PDF (for chats that read PDFs page by page), without LibreOffice."""
    from pptx import Presentation

    prs = Presentation(str(pptx_path))
    pages = [SlidePainter(prs, slide, long_edge).paint() for slide in prs.slides]
    if not pages:
        raise ValueError("This deck has no slides.")
    out_pdf.parent.mkdir(parents=True, exist_ok=True)
    pages[0].save(out_pdf, "PDF", resolution=150, save_all=True, append_images=pages[1:])
    return out_pdf
